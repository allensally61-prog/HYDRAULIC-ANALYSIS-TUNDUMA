#!/usr/bin/env python3
"""Excel zone calculator: one sheet per hydraulic zone, every result a live Excel formula.

    python make_zone_workbook.py   -> results/Tunduma_Zone_Calculator.xlsx

The starting values come from tunduma_config.json and the survey alignment (same model as
run_analysis.py). In Excel, change any blue input cell and the zone recalculates: flows,
Hazen-Williams losses, HGL, Bernoulli pump head / gravity surplus, power, motor, energy,
pressures at critical points, PN class and surge. The 'Profile' sheet checks the pressure at
every surveyed point.
"""
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import build_model
import water_hydraulics as wh

FONT = "Arial"
F_IN = Font(name=FONT, size=10, color="0000FF")             # inputs: blue
F_CALC = Font(name=FONT, size=10, color="000000")           # formulas: black
F_LINK = Font(name=FONT, size=10, color="008000")           # links to other sheets: green
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=10, italic=True, color="555555")
F_KEY = Font(name=FONT, size=11, bold=True)
FILL_IN = PatternFill("solid", fgColor="FFFF99")             # key assumptions / editable
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_RES = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class Sheet:
    """Small helper to write labelled rows and remember cell addresses."""

    def __init__(self, ws):
        self.ws = ws
        self.r = 1

    def cell(self, row, col, value, font=F_CALC, fmt=None, fill=None, border=True, align=None):
        c = self.ws.cell(row=row, column=col, value=value)
        c.font = font
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if border:
            c.border = BOX
        if align:
            c.alignment = align
        return c

    def title(self, text, sub=None):
        self.cell(self.r, 1, text, F_TITLE, border=False)
        self.r += 1
        if sub:
            self.cell(self.r, 1, sub, F_SUB, border=False)
            self.r += 1
        self.r += 1

    def section(self, text, width=8):
        for c in range(1, width + 1):
            self.cell(self.r, c, text if c == 1 else None, F_BOLD, fill=FILL_SEC)
        self.r += 1

    def row(self, label, value, unit="", note="", kind="calc", fmt="#,##0.00"):
        """Label | value | unit | note. kind: input, calc, link, key."""
        font = {"input": F_IN, "calc": F_CALC, "link": F_LINK, "key": F_KEY}[kind]
        self.cell(self.r, 1, label)
        c = self.cell(self.r, 2, value, font, fmt, FILL_IN if kind == "input" else
                      (FILL_RES if kind == "key" else None))
        self.cell(self.r, 3, unit)
        self.cell(self.r, 4, note, F_SUB, align=Alignment(wrap_text=False))
        ref = f"$B${self.r}"
        self.r += 1
        return ref

    def header(self, cols, row=None):
        row = row or self.r
        for i, h in enumerate(cols, 1):
            self.cell(row, i, h, F_HEAD, fill=FILL_HEAD, align=Alignment(wrap_text=True, vertical="center"))
        self.ws.row_dimensions[row].height = 42
        self.r = row + 1


def q(name):
    """Quoted sheet name for formulas."""
    return f"'{name}'"


def build(out_path, cfg_path="tunduma_config.json"):
    cfg = build_model.load_config(cfg_path)
    al = build_model.prepare_alignment(cfg)
    data, routes = build_model.build_engine_model(cfg, al)
    s = wh.System(data).solve()
    st = s.settings
    N, L = s.nodes, {l.id: l for l in s.links}
    fac = {f["id"]: f for f in cfg["facilities"]}
    lcfg = {l["id"]: l for l in cfg["links"]}
    wb = Workbook()

    # ================================================================ Read Me
    ws = wb.active
    ws.title = "Read Me"
    sh = Sheet(ws)
    sh.title("Tunduma Water Supply – Zone Calculator",
             "Hydraulic analysis from the Momba River intake to the Tunduma tanks, split by zone. "
             "Every result is a live Excel formula.")
    lines = [
        ("WHEN YOU RECEIVE THE REAL LEVELS – 4 STEPS", None),
        ("Step 1 – sheet 'Levels'", "type ground, floor, MINIMUM and MAXIMUM water level, inlet level and capacity of "
                                    "every tank / sump / river. Read the 'Check' column: it must say OK."),
        ("Step 2 – sheet 'Demands'", "type the daily demand of each tank and the WTP losses (if they changed)."),
        ("Step 3 – zone sheets Z1 … Z7", "only if pipes or pumps changed: DN, length, C, hours, number of pumps, "
                                       "efficiencies, ER duty (blue cells)."),
        ("Step 4 – sheet 'Summary'", "read the pump heads (at minimum and maximum suction level), gravity surplus, "
                                   "tank results and storage. Each zone sheet explains its result in section F."),
        ("Optional – Python and PDFs", "run 'python levels_from_excel.py results/Tunduma_Zone_Calculator.xlsx' to copy "
                                       "the levels into tunduma_config.json, then 'python run_analysis.py' and the PDF "
                                       "scripts to regenerate every report with the same levels."),
        ("", None),
        ("COLOUR LEGEND", None),
        ("Blue text, yellow fill", "input – edit these"),
        ("Black text", "formula – do not overwrite"),
        ("Green text", "link from another sheet"),
        ("Green fill, bold", "key result"),
        ("", None),
        ("WHY MINIMUM AND MAXIMUM LEVELS MATTER", None),
        ("Pumped zone, design case", "suction tank at its MINIMUM level → the pump must lift the most → TDH max "
                                     "(used to size pump and motor)."),
        ("Pumped zone, low-head case", "suction tank at its MAXIMUM level → least lift → TDH min (the pump must also "
                                       "run safely here: it gives more flow unless throttled / VFD)."),
        ("Gravity zone, design case", "supplying tank at its MINIMUM level → least energy available → must still "
                                      "reach every tank (surplus ≥ 0)."),
        ("Gravity zone, high case", "supplying tank at its MAXIMUM level → most energy → the valve must burn the "
                                    "biggest surplus; static pressure (pipe class) uses this level."),
        ("Receiving tank", "water must arrive at the inlet level (or the maximum level if the inlet is lower / "
                           "submerged) plus a small residual head."),
        ("", None),
        ("ZONES (split at every free water surface)", None),
        ("Z1 River-Intake", "gravity: Momba River → intake sump (abstraction pipe)"),
        ("Z2 Intake-Grit", "pumped: low-lift pumps, intake sump → grit tank"),
        ("Z3 Grit-WTP", "pumped: high-lift pumps, grit tank → WTP inlet (raw-water main)"),
        ("Z4 WTP-Ikana", "pumped: WTP clear-water pumps → Ikana GBR (treated-water rising main)"),
        ("Z5 Ikana-Nkangamo", "gravity: Ikana GBR → Nkangamo BPS sump"),
        ("Z6 Nkangamo-Uhuru", "pumped: Nkangamo BPS → J-001N (ESR), J-002N (GSR01), J-003N (Ipito) → Uhuru Park GBR"),
        ("Z7 Uhuru-Chapwa", "gravity: Uhuru Park GBR → J-004N → Makambini GSR03 and Chapwa GSR04"),
        ("", None),
        ("METHOD (Bernoulli / energy equation between two water surfaces)", None),
        ("z1 + p1/ρg + V1²/2g + Hp = z2 + p2/ρg + V2²/2g + hL", "at tank surfaces p = 0 and V ≈ 0"),
        ("Pumped zone", "Hp = H_req − z1 ;  TDH = Hp + station losses"),
        ("Gravity zone", "surplus = z1 − H_req  (≥ 0 OK, burnt by an FCV/PRV; < 0 short of head)"),
        ("H_req (head needed at zone start)", "largest of: each destination arrival level + residual + losses on the "
                                              "way, and each critical point crown + min. pressure + losses up to it"),
        ("Friction (Hazen-Williams)", "hf = 10.67 L Q^1.852 / (C^1.852 D^4.8704) ;  hL = hf × (1 + minor-loss factor)"),
        ("Pressure at a point", "p = HGL − pipe crown ;  working = max(p, static level − crown) ;  "
                                "PN = smallest class ≥ working × 0.0981 bar/m"),
        ("Surge (Joukowsky)", "Δh = a V / g"),
        ("", None),
        ("SOURCE OF THE STARTING VALUES", None),
        ("Employer's Requirements", "flows, pipe sizes and lengths, pump duties, tank elevations and capacities."),
        ("ASSUMED (to confirm)", "operating water levels, inlet levels, WTP inlet, C value, station losses, pipe "
                                 "cover – marked in 'Levels' column M and in cell notes."),
        ("Check", "column 'Python model' on Summary = run_analysis.py with the starting inputs."),
    ]
    for a, b in lines:
        bold = b is None and a.isupper()
        sh.cell(sh.r, 1, a, F_BOLD if bold else F_CALC, border=False)
        if b:
            sh.cell(sh.r, 2, b, F_CALC, border=False)
        if a == "Blue text, yellow fill":
            ws.cell(sh.r, 1).font, ws.cell(sh.r, 1).fill = F_IN, FILL_IN
        if a == "Green text":
            ws.cell(sh.r, 1).font = F_LINK
        if a == "Green fill, bold":
            ws.cell(sh.r, 1).font, ws.cell(sh.r, 1).fill = F_KEY, FILL_RES
        sh.r += 1
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 125

    # ================================================================ Global
    ws = wb.create_sheet("Global")
    sh = Sheet(ws)
    sh.title("Global constants and design criteria", "Used by every zone (named cells).")
    names = {}

    def gname(name, label, value, unit, note, fmt="0.00"):
        ref = sh.row(label, value, unit, note, "input", fmt)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Global!{ref}")
        names[name] = ref

    sh.section("Constants and criteria", 4)
    gname("g", "Gravitational acceleration g", 9.81, "m/s²", "physical constant")
    gname("k_minor", "Minor losses as fraction of friction", st["minor_loss_factor"], "–",
          "ASSUMED 10 % (bends, valves, fittings)")
    gname("h_station", "Losses inside each pump station", st["station_losses"], "m", "ASSUMED pipework/valves/meter")
    gname("h_res", "Residual head above a tank inlet", st["terminal_residual_head"], "m", "ASSUMED safety margin")
    gname("cover", "Pipe crown depth below ground", st["pipe_cover"], "m", "ASSUMED")
    gname("p_min", "Minimum pressure at pipe crown along mains", s.pipe_pmin, "m", "ASSUMED")
    gname("p_min_j", "Minimum pressure at junctions / delivery", st["min_pressure_head"], "m", "ER Table 7")
    gname("wave", "Pressure wave speed a (steel)", cfg["pipe_defaults"]["wave_speed"], "m/s", "ASSUMED, for surge")
    gname("motor_margin", "Motor size margin over shaft power", st["motor_margin"], "–", "common practice")
    gname("tariff", "Electricity tariff", st["electricity_tariff"], "per kWh", "enter the TANESCO tariff")
    gname("bar_per_m", "Bar per metre of water", 0.0981, "bar/m", "ρg/10⁵")
    gname("C_default", "Default Hazen-Williams C", cfg["pipe_defaults"]["hw_c"], "–", "ASSUMED, cement-lined steel (aged)")
    gname("fcv_limit", "Excess head above which a flow control valve is flagged", 5.0, "m", "judgement")
    sh.r += 1
    sh.section("Pipe pressure classes (PN, bar)", 4)
    start = sh.r
    for pn in st["pn_classes"]:
        sh.cell(sh.r, 1, f"PN{pn}")
        sh.cell(sh.r, 2, pn, F_IN, "0", FILL_IN)
        sh.r += 1
    wb.defined_names["PN_list"] = DefinedName("PN_list", attr_text=f"Global!$B${start}:$B${sh.r - 1}")
    sh.r += 1
    sh.section("Standard motor sizes (kW)", 4)
    start = sh.r
    for kw in wh.MOTOR_KW:
        sh.cell(sh.r, 2, kw, F_IN, "0.0#", FILL_IN)
        sh.r += 1
    wb.defined_names["Motor_list"] = DefinedName("Motor_list", attr_text=f"Global!$B${start}:$B${sh.r - 1}")
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 9
    ws.column_dimensions["D"].width = 48

    # ================================================================ Levels
    ws = wb.create_sheet("Levels")
    sh = Sheet(ws)
    sh.title("Levels of every water surface – ENTER THE REAL LEVELS HERE",
             "Minimum and maximum water levels drive every zone. Blank inlet = the maximum level is used. "
             "All levels in m above datum (survey datum).")
    sh.header(["ID", "Facility", "Role", "Ground level (m)", "Floor level (m)", "MINIMUM water level (m)",
               "MAXIMUM water level (m)", "Inlet level (m)", "Capacity (m³)", "Water depth (m)", "Plan area (m²)",
               "Arrival level used (m)", "Source / status", "Check"])
    roles = {"MOMBA_RIVER": "source: minimum = dry-season low level, maximum = flood level",
             "MOMBA_INTAKE": "intake sump: minimum = pump start/stop low level (submergence)",
             "GRIT_TANK": "grit tank = high-lift pump suction; receives low-lift flow",
             "MOMBA_WTP": "inlet = aerator inlet; min/max/capacity = CLEAR-WATER tank (pump suction)",
             "IKANA_GBR": "receives WTP pumping; supplies gravity main",
             "NKANGAMO_BPS": "BPS sump: receives gravity main; BPS pump suction",
             "UHURU_PARK_GBR": "receives BPS main; supplies town + J-004N",
             "TUNDUMA_ESR": "existing elevated tank (from J-001N)",
             "UHURU_GSR01": "existing tank (from J-002N)",
             "IPITO_GSR02": "existing tank (from J-003N)",
             "MAKAMBINI_GSR03": "existing tank (from J-004N, gravity)",
             "CHAPWA_GSR04": "existing tank (from J-004N, gravity)"}
    LV = {}
    for fid in roles:
        f = fac[fid]
        nd = N[fid]
        r = sh.r
        note = f.get("_note", "")
        status = ("levels ASSUMED – confirm" if "ASSUMED" in note or "CONFIRM" in note.upper() else "ER")
        if "ER:" in note and status != "ER":
            status = "ER elevation/capacity; " + status
        sh.cell(r, 1, fid, F_BOLD)
        sh.cell(r, 2, f.get("name", fid))
        sh.cell(r, 3, roles[fid], F_SUB)
        sh.cell(r, 4, nd.elevation, F_IN, "#,##0.00", FILL_IN)
        floor = f.get("floor")
        sh.cell(r, 5, floor, F_IN, "#,##0.00", FILL_IN)
        sh.cell(r, 6, nd.min_level, F_IN, "#,##0.00", FILL_IN)
        sh.cell(r, 7, nd.max_level, F_IN, "#,##0.00", FILL_IN)
        sh.cell(r, 8, f.get("inlet_level"), F_IN, "#,##0.00", FILL_IN)
        sh.cell(r, 9, nd.volume, F_IN, "#,##0", FILL_IN)
        sh.cell(r, 10, f'=IF(OR(E{r}="",G{r}=""),"",G{r}-E{r})', fmt="0.00")
        sh.cell(r, 11, f'=IF(OR(I{r}="",J{r}="",J{r}<=0),"",I{r}/J{r})', fmt="#,##0")
        if fid == "MOMBA_INTAKE":
            arr = f"=F{r}"
            cmt = "Intake: water enters below the surface, so the sump MINIMUM level is what the river must hold."
        elif fid == "MOMBA_WTP":
            arr = f'=IF(H{r}="",G{r},H{r})'
            cmt = "WTP: raw water must reach the aerator inlet (column H); min/max are the clear-water tank."
        else:
            arr = f'=IF(H{r}="",G{r},MAX(H{r},G{r}))'
            cmt = "Water must reach the inlet; if the inlet is below the maximum level (submerged) the maximum level counts."
        sh.cell(r, 12, arr, fmt="#,##0.00").comment = Comment(cmt, "model")
        sh.cell(r, 13, status, F_IN, fill=FILL_IN)
        sh.cell(r, 14, f'=IF(OR(F{r}="",G{r}=""),"ENTER MIN AND MAX",IF(G{r}<=F{r},"MAX must be above MIN",'
                       f'IF(AND(E{r}<>"",F{r}<E{r}),"MIN below floor",IF(AND(H{r}<>"",H{r}<G{r}),'
                       f'"inlet below MAX – MAX used","OK"))))')
        LV[fid] = {k: f"Levels!${c}${r}" for k, c in
                   (("ground", "D"), ("floor", "E"), ("min", "F"), ("max", "G"), ("inlet", "H"), ("cap", "I"),
                    ("arr", "L"))}
        sh.r += 1
    sh.r += 1
    sh.cell(sh.r, 1, "Notes", F_BOLD, border=False)
    for t in ["MINIMUM level = lowest operating level (pump cut-out / bottom water level, BWL). "
              "MAXIMUM level = top water level (TWL, overflow level).",
              "Floor and capacity are used for the depth/area check and the Storage sheet; they do not change the "
              "hydraulics.",
              "If the real inlet is a free fall above the tank, enter its level in 'Inlet level'; if it enters from "
              "below, leave it blank.",
              "After typing new levels, look at Summary – every zone, tank and pressure check updates."]:
        sh.r += 1
        sh.cell(sh.r, 1, t, F_SUB, border=False)
    for i, w in enumerate([18, 36, 58, 10, 10, 12, 12, 10, 10, 9, 10, 11, 40, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D5"

    # ================================================================ Demands
    ws = wb.create_sheet("Demands")
    sh = Sheet(ws)
    sh.title("Demands and design flows", "Daily volumes (m³/d) supplied by each tank; pipe volumes follow by mass balance.")
    sh.section("Tank demands (ER schematic, design year 2035)", 4)
    dem = {}
    for tid, label in (("UHURU_PARK_GBR", "Uhuru Park GBR – town supply"), ("UHURU_GSR01", "Existing Uhuru GSR 01"),
                       ("TUNDUMA_ESR", "Existing Tunduma ESR"), ("IPITO_GSR02", "Existing Ipito GSR 02"),
                       ("MAKAMBINI_GSR03", "Existing Makambini GSR 03"), ("CHAPWA_GSR04", "Existing Chapwa GSR 04")):
        dem[tid] = sh.row(label, N[tid].demand, "m³/d", "ER schematic", "input", "#,##0.00")
    loss = sh.row("WTP + intake-to-WTP losses", N["MOMBA_WTP"].loss, "–", "ER Table 8: 2 % + 5 %", "input", "0.0%")
    sh.r += 1
    sh.section("Daily volume in each pipe (mass balance, downstream → upstream)", 4)
    D = {}
    D["MAKAMBINI_BRANCH"] = sh.row("J-004N → Makambini", f"={dem['MAKAMBINI_GSR03']}", "m³/d")
    D["CHAPWA_BRANCH"] = sh.row("J-004N → Chapwa", f"={dem['CHAPWA_GSR04']}", "m³/d")
    D["UHURU_J004N"] = sh.row("Uhuru Park → J-004N", f"={D['MAKAMBINI_BRANCH']}+{D['CHAPWA_BRANCH']}", "m³/d")
    D["UHURU_INLET"] = sh.row("J-003N → Uhuru Park GBR", f"={dem['UHURU_PARK_GBR']}+{D['UHURU_J004N']}", "m³/d")
    D["IPITO_BRANCH"] = sh.row("J-003N → Ipito", f"={dem['IPITO_GSR02']}", "m³/d")
    D["GSR01_BRANCH"] = sh.row("J-002N → GSR01", f"={dem['UHURU_GSR01']}", "m³/d")
    D["ESR_BRANCH"] = sh.row("J-001N → ESR", f"={dem['TUNDUMA_ESR']}", "m³/d")
    D["BPS_RM_3"] = sh.row("J-002N → J-003N", f"={D['UHURU_INLET']}+{D['IPITO_BRANCH']}", "m³/d")
    D["BPS_RM_2"] = sh.row("J-001N → J-002N", f"={D['BPS_RM_3']}+{D['GSR01_BRANCH']}", "m³/d")
    D["BPS_RM_1"] = sh.row("Nkangamo BPS → J-001N", f"={D['BPS_RM_2']}+{D['ESR_BRANCH']}", "m³/d")
    D["IKANA_NKANGAMO_GM"] = sh.row("Ikana GBR → Nkangamo BPS", f"={D['BPS_RM_1']}", "m³/d")
    D["WTP_IKANA_RM"] = sh.row("WTP → Ikana GBR (net WTP output)", f"={D['IKANA_NKANGAMO_GM']}", "m³/d", "", "key")
    D["RAW_MAIN"] = sh.row("Raw water: intake → WTP", f"={D['WTP_IKANA_RM']}/(1-{loss})", "m³/d",
                           "net ÷ (1 − losses)", "key")
    D["LL_DELIVERY"] = sh.row("Low-lift → grit tank", f"={D['RAW_MAIN']}", "m³/d")
    D["ABSTRACTION"] = sh.row("River → intake sump", f"={D['RAW_MAIN']}", "m³/d")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 30

    # ================================================================ zones
    def points_for(lid, extra=()):
        """Critical points of a pipe from the model: governing high point and deepest point."""
        l = L[lid]
        if l.id not in routes:
            return list(extra)
        inner = [r for r in l.rows[1:-1]]
        if not inner:
            return list(extra)
        hi = max(inner, key=lambda r: r[1] + l.hloss * r[0] / l.length)
        lo = min(inner, key=lambda r: r[1])
        pts = [(lid, hi[0], hi[1] + st["pipe_cover"], "critical high point"),
               (lid, lo[0], lo[1] + st["pipe_cover"], "lowest point (highest static pressure)")]
        if l.gov_point:
            x, z = l.gov_point
            if all(abs(x - p[1]) > 1 for p in pts):
                pts.insert(0, (lid, x, z + st["pipe_cover"], "governing point (model)"))
        return pts + list(extra)

    zones = [
        dict(sheet="Z1 River-Intake", title="Zone 1 – Momba River to intake sump (gravity)", kind="gravity",
             start="MOMBA_RIVER", hours=N["MOMBA_INTAKE"].pumps["operating_hours"],
             hours_note="carries the intake pump flow while the pumps run",
             explain="The river fills the intake sump by gravity through the abstraction pipe. The river at its "
                     "MINIMUM (dry-season) level must stay above the sump MINIMUM level plus the pipe loss, otherwise "
                     "the low-lift pumps lose submergence.",
             pipes=[("ABSTRACTION", None)], dests=[("MOMBA_INTAKE", "ABSTRACTION", 0.0)], points=[]),
        dict(sheet="Z2 Intake-Grit", title="Zone 2 – Intake low-lift pumps to grit tank (pumped)", kind="pumped",
             start="MOMBA_INTAKE", hours=N["MOMBA_INTAKE"].pumps["operating_hours"],
             explain="Low-lift pumps lift raw water from the intake sump into the grit tank. Pump head = level the "
                     "water must reach in the grit tank + losses − sump level.",
             pipes=[("LL_DELIVERY", None)], dests=[("GRIT_TANK", "LL_DELIVERY", None)], points=[]),
        dict(sheet="Z3 Grit-WTP", title="Zone 3 – Grit tank high-lift pumps to the WTP (pumped)", kind="pumped",
             start="GRIT_TANK", hours=N["GRIT_TANK"].pumps["operating_hours"],
             explain="High-lift pumps push raw water from the grit tank through the raw-water main into the WTP "
                     "inlet (cascade aerator).",
             pipes=[("RAW_MAIN", None)], dests=[("MOMBA_WTP", "RAW_MAIN", None)], points=points_for("RAW_MAIN")),
        dict(sheet="Z4 WTP-Ikana", title="Zone 4 – WTP clear-water pumps to Ikana GBR (pumped)", kind="pumped",
             start="MOMBA_WTP", hours=N["MOMBA_WTP"].pumps["operating_hours"],
             explain="The WTP pumps lift treated water about 300 m over ~20 km from the clear-water tank to Ikana GBR "
                     "on the hill. Most of the head is lift; friction is the rest. Pressure is highest at the pumps.",
             pipes=[("WTP_IKANA_RM", None)], dests=[("IKANA_GBR", "WTP_IKANA_RM", None)],
             points=points_for("WTP_IKANA_RM")),
        dict(sheet="Z5 Ikana-Nkangamo", title="Zone 5 – Ikana GBR to Nkangamo BPS (gravity)", kind="gravity",
             start="IKANA_GBR", hours=st["gravity_hours"],
             explain="Water runs downhill by gravity from Ikana GBR to the Nkangamo sump. With Ikana at its MINIMUM "
                     "level it must still arrive (surplus ≥ 0); any surplus is burnt by the flow control valve at "
                     "the sump inlet. Static pressure uses Ikana at its MAXIMUM level.",
             pipes=[("IKANA_NKANGAMO_GM", None)], dests=[("NKANGAMO_BPS", "IKANA_NKANGAMO_GM", None)],
             points=points_for("IKANA_NKANGAMO_GM")),
        dict(sheet="Z6 Nkangamo-Uhuru", title="Zone 6 – Nkangamo BPS to Uhuru Park GBR and existing tanks (pumped)",
             kind="pumped", start="NKANGAMO_BPS", hours=N["NKANGAMO_BPS"].pumps["operating_hours"],
             explain="One pump station feeds four tanks (ESR, GSR01, Ipito, Uhuru Park) and crosses a high point near "
                     "74 km. Bernoulli is applied to each; the one needing the most head sets the pumps, the others "
                     "get flow control valves.",
             pipes=[("BPS_RM_1", None), ("BPS_RM_2", "BPS_RM_1"), ("BPS_RM_3", "BPS_RM_2"),
                    ("UHURU_INLET", "BPS_RM_3"), ("ESR_BRANCH", "BPS_RM_1"), ("GSR01_BRANCH", "BPS_RM_2"),
                    ("IPITO_BRANCH", "BPS_RM_3")],
             dests=[("TUNDUMA_ESR", "ESR_BRANCH", None), ("UHURU_GSR01", "GSR01_BRANCH", None),
                    ("IPITO_GSR02", "IPITO_BRANCH", None), ("UHURU_PARK_GBR", "UHURU_INLET", None)],
             points=points_for("BPS_RM_1") + points_for("IPITO_BRANCH")),
        dict(sheet="Z7 Uhuru-Chapwa", title="Zone 7 – Uhuru Park GBR to Makambini and Chapwa (gravity)",
             kind="gravity", start="UHURU_PARK_GBR", hours=st["gravity_hours"],
             explain="Uhuru Park GBR feeds Makambini and Chapwa by gravity through J-004N. Chapwa is far below, so it "
                     "arrives with large excess head (valve + higher PN). Check the points near the GBR: the ground "
                     "there is close to the tank level.",
             pipes=[("UHURU_J004N", None), ("MAKAMBINI_BRANCH", "UHURU_J004N"), ("CHAPWA_BRANCH", "UHURU_J004N")],
             dests=[("MAKAMBINI_GSR03", "MAKAMBINI_BRANCH", None), ("CHAPWA_GSR04", "CHAPWA_BRANCH", None)],
             points=points_for("UHURU_J004N") + points_for("CHAPWA_BRANCH")),
    ]

    prof_last = 4 + sum(len(l.points) for l in s.links)      # Profile sheet: header row 4, data from row 5
    zone_refs = {}
    pipe_rows = {}                     # pipe id -> (sheet, row)
    dest_rows = {}                     # tank id -> (sheet, row)
    for z in zones:
        ws = wb.create_sheet(z["sheet"])
        sh = Sheet(ws)
        sh.title(z["title"], z["explain"])
        ws.cell(2, 1).alignment = Alignment(wrap_text=False)
        lv = LV[z["start"]]
        sh.section("A. Inputs (levels come from sheet 'Levels')", 4)
        z1 = sh.row(f"Start: {z['start']} MINIMUM water level", f"={lv['min']}", "m a.s.l.",
                    "point 1 of Bernoulli, design case (least energy at the start)", "link")
        z1max = sh.row(f"Start: {z['start']} MAXIMUM water level", f"={lv['max']}", "m a.s.l.",
                       "low-head case for pumps / highest energy for gravity", "link")
        hours = sh.row("Operating hours per day", z["hours"], "h/d",
                       z.get("hours_note", "ER: pumps 22 h/d; gravity 24 h/d"), "input", "0.0")
        if z["kind"] == "gravity":
            static_f = f"={z1max}"
            static_note = "MAXIMUM level of the supplying tank (water at rest)"
        else:
            static_f = "=MAX(" + ",".join(LV[t]["arr"] for t, _, _ in z["dests"]) + ")"
            static_note = "highest arrival level at the end of the zone (pipe full, pumps stopped)"
        static = sh.row("Static level for pressure class", static_f, "m a.s.l.", static_note, "calc")
        if z["kind"] == "pumped":
            p = fac[z["start"]]["pumps"]
            duty = sh.row("Number of duty pumps", p["duty"], "no.", "working pumps sharing the flow", "input", "0")
            sh.row("Number of standby pumps", p["standby"], "no.", "for information", "input", "0")
            eff = sh.row("Pump efficiency η_p", p["efficiency"], "–", "ASSUMED – use the pump curve", "input", "0.00")
            meff = sh.row("Motor efficiency η_m", p["motor_efficiency"], "–", "ASSUMED", "input", "0.00")
            sh.row("ER duty flow per pump", p["er_duty"]["q_m3h"], "m³/h", "Employer's Requirements", "input", "#,##0")
            er_h = sh.row("ER duty head", p["er_duty"]["head_m"], "m", "Employer's Requirements", "input", "#,##0.0")
        sh.r += 1

        # ---- pipe table
        sh.section("B. Pipes – flow (continuity) and friction (Hazen-Williams)", 21)
        cols = ["Pipe", "From", "To", "Upstream pipe", "DN", "Internal D (mm)", "Length L (m)", "C (H-W)",
                "Daily volume (m³/d)", "Q (m³/s)", "Q (l/s)", "Area (m²)", "V (m/s)", "hf friction (m)",
                "hm minor (m)", "hL total (m)", "Gradient (m/km)", "Loss before this pipe (m)", "HGL start (m)",
                "HGL end (m)", "Surge Δh = aV/g (m)"]
        sh.header(cols)
        first = sh.r
        rows = {}
        for lid, parent in z["pipes"]:
            l = L[lid]
            r = sh.r
            rows[lid] = r
            pipe_rows[lid] = (z["sheet"], r)
            sh.cell(r, 1, lid, F_BOLD)
            sh.cell(r, 2, l.frm)
            sh.cell(r, 3, l.to)
            sh.cell(r, 4, parent or "–")
            sh.cell(r, 5, l.dn, F_IN, "0", FILL_IN)
            sh.cell(r, 6, f"=E{r}", F_CALC, "0.0").comment = Comment(
                "Internal diameter. Defaults to DN; overwrite with the real ID (e.g. lined steel) if known.", "model")
            sh.cell(r, 7, round(l.length, 1), F_IN, "#,##0.0", FILL_IN).comment = Comment(
                lcfg[lid].get("er", "") + (f" | ER length {lcfg[lid]['er_length']:,} m" if lcfg[lid].get('er_length')
                                           else "") + " | model length measured along the surveyed slope", "model")
            sh.cell(r, 8, "=C_default", F_IN, "0", FILL_IN)
            sh.cell(r, 9, f"=Demands!{D[lid]}", F_LINK, "#,##0.0")
            sh.cell(r, 10, f"=I{r}/({hours}*3600)", fmt="0.0000")
            sh.cell(r, 11, f"=J{r}*1000", fmt="0.0")
            sh.cell(r, 12, f"=PI()*(F{r}/1000)^2/4", fmt="0.0000")
            sh.cell(r, 13, f"=J{r}/L{r}", fmt="0.000")
            sh.cell(r, 14, f"=10.67*G{r}*J{r}^1.852/(H{r}^1.852*(F{r}/1000)^4.8704)", fmt="0.00")
            sh.cell(r, 15, f"=k_minor*N{r}", fmt="0.00")
            sh.cell(r, 16, f"=N{r}+O{r}", fmt="0.00")
            sh.cell(r, 17, f"=P{r}/G{r}*1000", fmt="0.00")
            sh.cell(r, 18, f"=R{rows[parent]}+P{rows[parent]}" if parent else 0, fmt="0.00")
            sh.cell(r, 21, f"=wave*M{r}/g", fmt="0.0")
            sh.r += 1
        last = sh.r - 1
        sh.r += 1

        # ---- destinations
        sh.section("C. Destinations – Bernoulli to each receiving water surface (levels from 'Levels')", 12)
        sh.header(["Destination", "Ends with pipe", "MIN level (m)", "MAX level (m)", "Arrival level (m)",
                   "Residual (m)", "z2 = arrival + residual (m)", "Head needed at zone start = z2 + losses (m)",
                   "Arriving HGL, design case (m)", "Excess head, design case (m)",
                   "Excess head, start at MAX level (m)", "Comment"])
        dfirst = sh.r
        for tid, lid, res in z["dests"]:
            r = sh.r
            pr = rows[lid]
            dest_rows[tid] = (z["sheet"], r)
            sh.cell(r, 1, tid, F_BOLD)
            sh.cell(r, 2, lid)
            sh.cell(r, 3, f"={LV[tid]['min']}", F_LINK, "#,##0.00")
            sh.cell(r, 4, f"={LV[tid]['max']}", F_LINK, "#,##0.00")
            sh.cell(r, 5, f"={LV[tid]['arr']}", F_LINK, "#,##0.00")
            sh.cell(r, 6, "=h_res" if res is None else res, F_IN, "0.00", FILL_IN)
            sh.cell(r, 7, f"=E{r}+F{r}", fmt="#,##0.00")
            sh.cell(r, 8, f"=G{r}+R{pr}+P{pr}", fmt="#,##0.00")
            sh.cell(r, 9, f"=T{pr}", fmt="#,##0.00")
            sh.cell(r, 10, f"=I{r}-G{r}", fmt="#,##0.00")
            sh.cell(r, 11, f"=J{r}+({z1max}-{z1})", fmt="#,##0.00")
            sh.cell(r, 12, f'=IF(J{r}<-0.005,"NOT REACHED – more head needed",IF(K{r}>fcv_limit,'
                           f'"fit flow control valve / PRV (burns "&TEXT(J{r},"0.0")&"–"&TEXT(K{r},"0.0")&" m)",'
                           f'"OK"))')
            sh.r += 1
        dlast = sh.r - 1
        sh.r += 1

        # ---- junctions (ER: minimum pressure at junctions / delivery points)
        juncs = [(lid_to, lid) for lid, _ in z["pipes"] for lid_to in [L[lid].to] if N[lid_to].type == "junction"]
        jfirst = jlast = None
        if juncs:
            sh.section("C2. Junctions – minimum pressure at offtakes (ER Table 7)", 8)
            sh.header(["Junction", "End of pipe", "Ground level (m)", "Minimum pressure (m)",
                       "Head needed at zone start = ground + p_min + losses (m)", "HGL at junction (m)",
                       "Pressure (m)", "Status"])
            jfirst = sh.r
            for jid, lid in juncs:
                r = sh.r
                pr = rows[lid]
                sh.cell(r, 1, jid, F_BOLD)
                sh.cell(r, 2, lid)
                sh.cell(r, 3, round(N[jid].elevation, 2), F_IN, "#,##0.00", FILL_IN)
                sh.cell(r, 4, "=p_min_j", F_IN, "0.0", FILL_IN)
                sh.cell(r, 5, f"=C{r}+D{r}+R{pr}+P{pr}", fmt="#,##0.00")
                sh.cell(r, 6, f"=T{pr}", fmt="#,##0.00")
                sh.cell(r, 7, f"=F{r}-C{r}", fmt="#,##0.0")
                sh.cell(r, 8, f'=IF(G{r}<D{r}-0.005,"below minimum","OK")')
                sh.r += 1
            jlast = sh.r - 1
            sh.r += 1

        # ---- critical points
        sh.section("D. Critical points along the pipes – pressure check (add your own rows)", 14)
        sh.header(["Pipe", "Chainage in pipe x (m)", "Ground level (m)", "Pipe crown (m)", "Head needed at zone "
                   "start (m)", "HGL at point (m)", "Pressure (m)", "Static (m)", "Working (m)", "Working (bar)",
                   "PN required", "Transient min = p − Δh (m)", "Note", "Status"])
        pfirst = sh.r
        n_blank = 3
        for i in range(len(z["points"]) + n_blank):
            r = sh.r
            if i < len(z["points"]):
                lid, x, ground, note = z["points"][i]
                sh.cell(r, 1, lid, F_IN, fill=FILL_IN)
                sh.cell(r, 2, round(x, 1), F_IN, "#,##0.0", FILL_IN)
                sh.cell(r, 3, round(ground, 2), F_IN, "#,##0.00", FILL_IN)
                sh.cell(r, 13, note, F_SUB)
            else:
                for c in (1, 2, 3):
                    sh.cell(r, c, None, F_IN, fill=FILL_IN)
                sh.cell(r, 13, "add a point: pipe ID, chainage, ground level", F_SUB)
            m = f"MATCH($A{r},$A${first}:$A${last},0)"
            idx = lambda col: f"INDEX(${col}${first}:${col}${last},{m})"   # noqa: E731
            blank = f'IF($A{r}="","",'
            sh.cell(r, 4, f"={blank}C{r}-cover)", fmt="#,##0.00")
            sh.cell(r, 5, f"={blank}D{r}+p_min+{idx('R')}+{idx('P')}*B{r}/{idx('G')})", fmt="#,##0.00")
            sh.cell(r, 6, f"={blank}{idx('S')}-{idx('P')}*B{r}/{idx('G')})", fmt="#,##0.00")
            sh.cell(r, 7, f"={blank}F{r}-D{r})", fmt="#,##0.0")
            sh.cell(r, 8, f"={blank}{static}-D{r})", fmt="#,##0.0")
            sh.cell(r, 9, f"={blank}MAX(G{r},H{r}))", fmt="#,##0.0")
            sh.cell(r, 10, f"={blank}I{r}*bar_per_m)", fmt="0.0")
            sh.cell(r, 11, f'={blank}IFERROR("PN"&SMALL(PN_list,COUNTIF(PN_list,"<"&J{r})+1),"> max PN"))')
            sh.cell(r, 12, f"={blank}G{r}-{idx('U')})", fmt="#,##0.0")
            sh.cell(r, 14, f'={blank}IF(G{r}<0,"NEGATIVE PRESSURE",IF(G{r}<p_min-0.005,"below minimum","OK")))')
            sh.r += 1
        plast = sh.r - 1
        sh.r += 1

        # ---- result
        sh.section("E. Bernoulli result", 4)
        dr_, pr_ = f"$H${dfirst}:$H${dlast}", f"$E${pfirst}:$E${plast}"
        jr_ = f"$E${jfirst}:$E${jlast}" if juncs else None
        hreq = sh.row("Head needed at zone start H_req", "=MAX(" + ",".join(x for x in (dr_, pr_, jr_) if x) + ")",
                      "m a.s.l.", "largest of sections C, C2 (junctions) and D (points)", "key")
        pick = (f'IF(MAX({pr_})>=MAX({dr_}),"point on "&INDEX($A${pfirst}:$A${plast},MATCH(MAX({pr_}),{pr_},0)),'
                f'INDEX($A${dfirst}:$A${dlast},MATCH(MAX({dr_}),{dr_},0)))')
        if juncs:
            pick = (f'IF(AND(MAX({jr_})>MAX({dr_}),MAX({jr_})>MAX({pr_})),"junction "&INDEX($A${jfirst}:$A${jlast},'
                    f'MATCH(MAX({jr_}),{jr_},0))&" (minimum pressure)",{pick})')
        gov = sh.row("Set by", "=" + pick, "", "governing destination, junction or point", "calc", "@")
        res = {"sheet": z["sheet"], "kind": z["kind"], "hreq": hreq, "gov": gov, "first": first, "z1": z1,
               "z1max": z1max, "hours": hours, "static": static}
        if z["kind"] == "pumped":
            hstart = hreq
            lift = sh.row("Hp = H_req − z1  (suction at MINIMUM level)", f"={hreq}-{z1}", "m",
                          "z1 + 0 + 0 + Hp = H_req  (p = 0, V ≈ 0 at the sump surface)")
            tdh = sh.row("TDH design (suction at MINIMUM level)", f"={lift}+h_station", "m",
                         "Hp + station losses → size pump and motor with this", "key")
            tdh_min = sh.row("TDH low-head (suction at MAXIMUM level)", f"={hreq}-{z1max}+h_station", "m",
                             "pump must also operate here (more flow unless throttled / VFD)", "key")
            qtot = sh.row("Total pumped flow", f"=J{first}", "m³/s", "flow in the first pipe", "calc", "0.0000")
            qp = sh.row("Flow per pump", f"={qtot}/{duty}", "m³/s", "", "calc", "0.0000")
            qph = sh.row("Flow per pump", f"={qp}*3600", "m³/h", "compare ER duty flow", "calc", "#,##0")
            ph = sh.row("Hydraulic power per pump P_h = ρgQH", f"=1000*g*{qp}*{tdh}/1000", "kW", "", "calc", "#,##0.0")
            ps = sh.row("Shaft power per pump P_s = P_h / η_p", f"={ph}/{eff}", "kW", "", "key", "#,##0.0")
            mreq = sh.row("Motor needed = P_s × margin", f"={ps}*motor_margin", "kW", "", "calc", "#,##0.0")
            motor = sh.row("Standard motor per pump", f'=IFERROR(SMALL(Motor_list,COUNTIF(Motor_list,"<"&{mreq})+1),'
                                                       f'"> list")', "kW", "next size up", "key", "#,##0.0")
            energy = sh.row("Electricity per day", f"={duty}*{ps}*{hours}/{meff}", "kWh/d", "", "calc", "#,##0")
            spec = sh.row("Specific energy", f"={energy}/I{first}", "kWh/m³", "", "calc", "0.000")
            sh.row("Electricity cost per year", f"={energy}*365*tariff", "per year", "uses Global tariff",
                   "calc", "#,##0")
            diff = sh.row("TDH design minus ER duty head", f"={tdh}-{er_h}", "m",
                          "negative = ER pump has spare head (throttle / VFD)", "calc", "+0.0;-0.0")
            res.update(tdh=tdh, tdh_min=tdh_min, motor=motor, energy=energy, ps=ps, er_h=er_h, diff=diff, qp=qp,
                       spec=spec, qph=qph)
        else:
            hstart = z1
            surplus = sh.row("Surplus = z1 − H_req  (start at MINIMUM level)", f"={z1}-{hreq}", "m",
                             "design case: must be ≥ 0", "key")
            surplus_max = sh.row("Surplus with start at MAXIMUM level", f"={z1max}-{hreq}", "m",
                                 "largest head the valve must burn", "key")
            sh.row("Status", f'=IF({surplus}<-0.005,"SHORT OF HEAD at minimum level – bigger pipe, lower outlet, '
                             f'higher tank or pump","OK – valve burns "&TEXT({surplus},"0.0")&" to "&'
                             f'TEXT({surplus_max},"0.0")&" m")', "", "", "calc", "@")
            res.update(surplus=surplus, surplus_max=surplus_max)
        for lid, parent in z["pipes"]:
            r = rows[lid]
            sh.cell(r, 19, f"={hstart}-R{r}", fmt="#,##0.00")
            sh.cell(r, 20, f"=S{r}-P{r}", fmt="#,##0.00")
        lowp = sh.row("Lowest pressure at any survey point (sheet Profile)",
                      f'=IF(COUNTIFS(Profile!$A$5:$A${prof_last},"{z["sheet"]}",Profile!$L$5:$L${prof_last},"yes")=0,'
                      f'"no points between tanks",_xlfn.MINIFS(Profile!$H$5:$H${prof_last},Profile!$A$5:$A${prof_last},'
                      f'"{z["sheet"]}",Profile!$L$5:$L${prof_last},"yes"))', "m",
                      "if below the minimum, add that point to section D", "calc", "#,##0.0")
        res["lowp"] = lowp
        sh.r += 1

        # ---- step-by-step explanation in words (live)
        sh.section("F. Step by step in words (updates with the numbers)", 4)
        crit_p = f"INDEX($A${dfirst}:$A${dlast},MATCH(MAX($H${dfirst}:$H${dlast}),$H${dfirst}:$H${dlast},0))"
        steps = [
            f'="1. Flow: "&TEXT(I{first},"#,##0")&" m³/d in "&TEXT({hours},"0")&" h → Q = "&TEXT(K{first},"0.0")'
            f'&" l/s, V = "&TEXT(M{first},"0.00")&" m/s in "&A{first}&"."',
            f'="2. Friction (Hazen-Williams): "&A{first}&" loses hL = "&TEXT(P{first},"0.00")&" m over "'
            f'&TEXT(G{first},"#,##0")&" m ("&TEXT(Q{first},"0.00")&" m/km)."',
            f'="3. Point 1 = {z["start"]} water surface at its MINIMUM level z1 = "&TEXT({z1},"0.00")&" m '
            f'(p = 0, V ≈ 0)."',
            f'="4. Point 2: the hardest place to reach is "&{gov}&"; it needs the HGL at the zone start to be '
            f'H_req = "&TEXT({hreq},"0.00")&" m (its level + residual + losses on the way)."',
        ]
        if z["kind"] == "pumped":
            steps += [
                f'="5. Bernoulli: z1 + Hp = H_req → Hp = "&TEXT({hreq},"0.00")&" − "&TEXT({z1},"0.00")&" = "'
                f'&TEXT({lift},"0.00")&" m; + station losses "&TEXT(h_station,"0.0")&" m → TDH = "'
                f'&TEXT({tdh},"0.0")&" m."',
                f'="6. With the suction tank full (MAXIMUM level "&TEXT({z1max},"0.00")&" m) TDH falls to "'
                f'&TEXT({tdh_min},"0.0")&" m: the pump must work between "&TEXT({tdh_min},"0.0")&" and "'
                f'&TEXT({tdh},"0.0")&" m at "&TEXT({qph},"#,##0")&" m³/h per pump."',
                f'="7. Power: P = ρgQH/η = "&TEXT({ps},"#,##0")&" kW per pump → motor "&TEXT({motor},"#,##0")'
                f'&" kW; ER duty head "&TEXT({er_h},"0")&" m (difference "&TEXT({diff},"+0.0;-0.0")&" m)."',
            ]
        else:
            steps += [
                f'="5. Bernoulli without pump: z1 = H_req + surplus → surplus = "&TEXT({z1},"0.00")&" − "'
                f'&TEXT({hreq},"0.00")&" = "&TEXT({surplus},"0.00")&" m."',
                f'="6. With the supplying tank full (MAXIMUM level "&TEXT({z1max},"0.00")&" m) the surplus is "'
                f'&TEXT({surplus_max},"0.00")&" m. "&IF({surplus}<-0.005,"Negative at minimum level: the zone '
                f'cannot deliver – change pipe or levels.","A flow control valve burns this so the flow stays at '
                f'design.")',
            ]
        steps += [
            f'="8. Pressure: lowest pressure along the route = "&TEXT({lowp},"0.0")&" m (minimum allowed "'
            f'&TEXT(p_min,"0")&" m); pipe class from section D and sheet Profile."',
        ]
        for t in steps:
            c = sh.cell(sh.r, 1, t, F_CALC, border=False)
            sh.r += 1
        res["hstart"] = hstart
        zone_refs[z["sheet"]] = res

        widths = [44, 14, 16, 16, 11, 11, 12, 13, 12, 12, 12, 30, 9, 11, 10, 10, 11, 12, 12, 12, 12]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B4"

    # ================================================================ Storage
    ws = wb.create_sheet("Storage")
    sh = Sheet(ws)
    sh.title("Storage – balancing for pump rest hours",
             "V_bal = inflow × (1 − hours/24); emergency = inflow × hours/24. Volumes and levels from 'Levels'.")
    sh.header(["Tank", "MIN level (m)", "MAX level (m)", "Daily inflow (m³/d)", "Inflow hours (h/d)",
               "Balancing volume (m³)", "Emergency hours", "Emergency volume (m³)", "Required (m³)",
               "Capacity (m³)", "Storage (h of flow)", "Status"])
    tank_in = {"IKANA_GBR": "WTP_IKANA_RM", "NKANGAMO_BPS": "IKANA_NKANGAMO_GM", "UHURU_PARK_GBR": "UHURU_INLET",
               "TUNDUMA_ESR": "ESR_BRANCH", "UHURU_GSR01": "GSR01_BRANCH", "IPITO_GSR02": "IPITO_BRANCH",
               "MAKAMBINI_GSR03": "MAKAMBINI_BRANCH", "CHAPWA_GSR04": "CHAPWA_BRANCH"}
    storage_rows = {}
    for tid, lid in tank_in.items():
        r = sh.r
        storage_rows[tid] = r
        sheet, pr = pipe_rows[lid]
        sn = q(sheet)
        sh.cell(r, 1, tid, F_BOLD)
        sh.cell(r, 2, f"={LV[tid]['min']}", F_LINK, "#,##0.00")
        sh.cell(r, 3, f"={LV[tid]['max']}", F_LINK, "#,##0.00")
        sh.cell(r, 4, f"={sn}!I{pr}", F_LINK, "#,##0")
        if tid == "NKANGAMO_BPS":
            bps_h = f"{q('Z6 Nkangamo-Uhuru')}!{zone_refs['Z6 Nkangamo-Uhuru']['hours']}"
            sh.cell(r, 5, f"={bps_h}", F_LINK, "0.0").comment = Comment(
                "Sump receives 24 h gravity flow and is emptied by the BPS pumps in these hours: it must hold the "
                "inflow during the pump-off period.", "model")
        else:
            sh.cell(r, 5, f"={sn}!I{pr}/({sn}!J{pr}*3600)", F_LINK, "0.0")
        sh.cell(r, 6, f"=D{r}*(1-E{r}/24)", fmt="#,##0")
        sh.cell(r, 7, 0, F_IN, "0.0", FILL_IN)
        sh.cell(r, 8, f"=D{r}*G{r}/24", fmt="#,##0")
        sh.cell(r, 9, f"=F{r}+H{r}", fmt="#,##0")
        sh.cell(r, 10, f'=IF({LV[tid]["cap"]}="","",{LV[tid]["cap"]})', F_LINK, "#,##0")
        sh.cell(r, 11, f'=IF(J{r}="","",J{r}/(D{r}/24))', fmt="0.0")
        sh.cell(r, 12, f'=IF(J{r}="","enter capacity",IF(J{r}>=I{r},"OK","TOO SMALL"))')
        sh.r += 1
    for i, w in enumerate([22, 11, 11, 13, 11, 13, 11, 13, 12, 12, 12, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ================================================================ Summary
    ws = wb.create_sheet("Summary", 1)
    sh = Sheet(ws)
    sh.title("Summary of all zones", "Linked to the zone sheets (green). Change levels on 'Levels' and read the "
                                     "result here. 'Python model' = run_analysis.py with the starting inputs.")
    sh.section("Zones", 17)
    sh.header(["Zone", "Type", "Flow (l/s)", "Velocity first pipe (m/s)", "hL first pipe (m)", "H_req (m a.s.l.)",
               "Set by", "Design: TDH (m) / surplus (m)", "Start at MAX level: TDH (m) / surplus (m)",
               "ER duty head (m)", "TDH − ER (m)", "Shaft kW per pump", "Motor kW per pump", "Energy (kWh/d)",
               "Lowest pressure on route (m)", "Python model (m)", "Excel − Python (m)"])
    py = {"Z1 River-Intake": N["MOMBA_RIVER"].surplus, "Z2 Intake-Grit": N["MOMBA_INTAKE"].station["tdh"],
          "Z3 Grit-WTP": N["GRIT_TANK"].station["tdh"], "Z4 WTP-Ikana": N["MOMBA_WTP"].station["tdh"],
          "Z5 Ikana-Nkangamo": N["IKANA_GBR"].surplus, "Z6 Nkangamo-Uhuru": N["NKANGAMO_BPS"].station["tdh"],
          "Z7 Uhuru-Chapwa": N["UHURU_PARK_GBR"].surplus}
    zfirst = sh.r
    for name, rf in zone_refs.items():
        r = sh.r
        sn = q(name)
        sh.cell(r, 1, name, F_BOLD)
        sh.cell(r, 2, rf["kind"])
        sh.cell(r, 3, f"={sn}!K{rf['first']}", F_LINK, "#,##0.0")
        sh.cell(r, 4, f"={sn}!M{rf['first']}", F_LINK, "0.00")
        sh.cell(r, 5, f"={sn}!P{rf['first']}", F_LINK, "#,##0.00")
        sh.cell(r, 6, f"={sn}!{rf['hreq']}", F_LINK, "#,##0.00")
        sh.cell(r, 7, f"={sn}!{rf['gov']}", F_LINK)
        main = rf.get("tdh", rf.get("surplus"))
        other = rf.get("tdh_min", rf.get("surplus_max"))
        sh.cell(r, 8, f"={sn}!{main}", F_LINK, "#,##0.0", FILL_RES)
        sh.cell(r, 9, f"={sn}!{other}", F_LINK, "#,##0.0")
        if rf["kind"] == "pumped":
            sh.cell(r, 10, f"={sn}!{rf['er_h']}", F_LINK, "#,##0.0")
            sh.cell(r, 11, f"={sn}!{rf['diff']}", F_LINK, "+0.0;-0.0")
            sh.cell(r, 12, f"={sn}!{rf['ps']}", F_LINK, "#,##0.0")
            sh.cell(r, 13, f"={sn}!{rf['motor']}", F_LINK, "#,##0.0")
            sh.cell(r, 14, f"={sn}!{rf['energy']}", F_LINK, "#,##0")
        else:
            for c in range(10, 15):
                sh.cell(r, c, "–")
        sh.cell(r, 15, f"={sn}!{rf['lowp']}", F_LINK, "#,##0.0")
        sh.cell(r, 16, round(py[name], 2), F_IN, "#,##0.00").comment = Comment(
            "Value from run_analysis.py (Python model) with the starting inputs. Fixed number, for checking.", "model")
        sh.cell(r, 17, f"=H{r}-P{r}", fmt="+0.00;-0.00;0.00")
        sh.r += 1
    tot = sh.r
    sh.cell(tot, 1, "Total pumping energy", F_BOLD)
    sh.cell(tot, 14, f"=SUM(N{zfirst}:N{tot - 1})", F_KEY, "#,##0", FILL_RES)
    sh.r += 1
    sh.cell(sh.r, 1, "Pumped zones: H = pump head with suction at MINIMUM level (design), I = with suction at "
                     "MAXIMUM level. Gravity zones: H/I = spare head with the supplying tank at MIN/MAX level "
                     "(≥ 0 OK, burnt by a valve; < 0 short of head).", F_SUB, border=False)
    sh.r += 2

    sh.section("Tanks and sumps", 12)
    sh.header(["Tank", "MIN level (m)", "MAX level (m)", "Arrival level (m)", "Fed by zone", "Arriving HGL (m)",
               "Excess head, design (m)", "Excess head, max (m)", "Inlet comment", "Capacity (m³)",
               "Storage required (m³)", "Storage status"])
    for tid in ["MOMBA_INTAKE", "GRIT_TANK", "MOMBA_WTP", "IKANA_GBR", "NKANGAMO_BPS", "TUNDUMA_ESR",
                "UHURU_GSR01", "IPITO_GSR02", "UHURU_PARK_GBR", "MAKAMBINI_GSR03", "CHAPWA_GSR04"]:
        r = sh.r
        sheet, dr = dest_rows[tid]
        sn = q(sheet)
        sh.cell(r, 1, tid, F_BOLD)
        sh.cell(r, 2, f"={LV[tid]['min']}", F_LINK, "#,##0.00")
        sh.cell(r, 3, f"={LV[tid]['max']}", F_LINK, "#,##0.00")
        sh.cell(r, 4, f"={LV[tid]['arr']}", F_LINK, "#,##0.00")
        sh.cell(r, 5, sheet)
        sh.cell(r, 6, f"={sn}!I{dr}", F_LINK, "#,##0.00")
        sh.cell(r, 7, f"={sn}!J{dr}", F_LINK, "#,##0.00")
        sh.cell(r, 8, f"={sn}!K{dr}", F_LINK, "#,##0.00")
        sh.cell(r, 9, f"={sn}!L{dr}", F_LINK)
        sh.cell(r, 10, f'=IF({LV[tid]["cap"]}="","",{LV[tid]["cap"]})', F_LINK, "#,##0")
        if tid in storage_rows:
            sh.cell(r, 11, f"=Storage!I{storage_rows[tid]}", F_LINK, "#,##0")
            sh.cell(r, 12, f"=Storage!L{storage_rows[tid]}", F_LINK)
        else:
            sh.cell(r, 11, "–")
            sh.cell(r, 12, "–")
        sh.r += 1
    sh.r += 1
    sh.cell(sh.r, 1, "Levels check (from sheet Levels)", F_BOLD, border=False)
    sh.r += 1
    sh.cell(sh.r, 1, '=IF(COUNTIF(Levels!$N$5:$N$16,"OK")=12,"All levels entered and consistent",'
                     '"CHECK sheet Levels – "&(12-COUNTIF(Levels!$N$5:$N$16,"OK"))&" row(s) need attention")',
            F_KEY, border=False)
    for i, w in enumerate([22, 10, 11, 11, 18, 26, 13, 13, 40, 11, 11, 13, 13, 12, 12, 12, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ================================================================ Profile
    ws = wb.create_sheet("Profile")
    sh = Sheet(ws)
    sh.title("Pressure at every surveyed point", "HGL from the zone sheets; ground from the topographic survey. "
                                                 "Column L = 'no' for pipe ends at a water surface (not checked).")
    sh.header(["Zone", "Pipe", "Chainage in pipe (m)", "Ground (m)", "Pipe crown (m)", "HGL (m)", "Static (m)",
               "Pressure (m)", "Working (m)", "PN required", "Status", "Check?"])
    static_of = {name: rf["static"] for name, rf in zone_refs.items()}
    for l in s.links:
        sheet, pr = pipe_rows[l.id]
        sn = q(sheet)
        n_pts = len(l.points)
        frm_free, to_free = s.nodes[l.frm].is_free_surface, s.nodes[l.to].is_free_surface
        for i, (x, zc) in enumerate(l.points):
            r = sh.r
            sh.cell(r, 1, sheet, border=False)
            sh.cell(r, 2, l.id, border=False)
            sh.cell(r, 3, round(x, 2), F_IN, "#,##0.0", border=False)
            sh.cell(r, 4, round(zc + st["pipe_cover"], 2), F_IN, "#,##0.00", border=False)
            sh.cell(r, 5, f"=D{r}-cover", fmt="#,##0.00", border=False)
            sh.cell(r, 6, f"={sn}!S{pr}-{sn}!P{pr}*C{r}/{sn}!G{pr}", F_LINK, "#,##0.00", border=False)
            sh.cell(r, 7, f"={sn}!{static_of[sheet]}-E{r}", F_LINK, "#,##0.0", border=False)
            sh.cell(r, 8, f"=F{r}-E{r}", fmt="#,##0.0", border=False)
            sh.cell(r, 9, f"=MAX(H{r},G{r})", fmt="#,##0.0", border=False)
            sh.cell(r, 10, f'=IFERROR("PN"&SMALL(PN_list,COUNTIF(PN_list,"<"&I{r}*bar_per_m)+1),"> max PN")',
                    border=False)
            check = not ((i == 0 and frm_free) or (i == n_pts - 1 and to_free))
            sh.cell(r, 11, f'=IF(L{r}="no","end at tank",IF(H{r}<0,"NEGATIVE",IF(H{r}<p_min-0.005,"low","OK")))',
                    border=False)
            sh.cell(r, 12, "yes" if check else "no", border=False)
            sh.r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:L{sh.r - 1}"
    for i, w in enumerate([20, 18, 12, 11, 11, 11, 10, 11, 11, 11, 12, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    order = ["Read Me", "Levels", "Summary", "Demands", "Global"] + [z["sheet"] for z in zones] + \
            ["Storage", "Profile"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 1
    for wsx in wb.worksheets:
        wsx.sheet_view.showGridLines = wsx.title in ("Profile",)
        wsx.sheet_view.tabSelected = wsx.title == "Levels"
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    os.makedirs("results", exist_ok=True)
    print(build(os.path.join("results", "Tunduma_Zone_Calculator.xlsx")))
