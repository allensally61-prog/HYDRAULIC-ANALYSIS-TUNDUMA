#!/usr/bin/env python3
"""Hydraulic analysis of a bulk water supply system.

    SOURCE -> INTAKE (+ raw water pumps) -> WTP (+ high-lift pumps)
           -> BOOSTER PUMP STATION(S) -> JUNCTIONS -> STORAGE TANKS

The system is described in a JSON file as nodes (source, intake, wtp, bps,
junction, tank) and pipes (links) forming a tree that starts at the source.
For every pipe and station the script works out:

  * design flows from the tank demands (mass balance, WTP process losses,
    pump operating hours)
  * velocity and friction losses (Hazen-Williams or Darcy-Weisbach) plus
    minor losses, with automatic DN sizing where "dn" is "auto"
  * the hydraulic grade line (HGL) and pressures along each pipe profile
  * pump duty point (Q, total dynamic head), power, motor size, energy and
    NPSH available
  * gravity mains: available vs. required head, excess head to throttle
  * surge (Joukowsky) and the PN class required along each pipe (EN 805)
  * sump retention and tank balancing storage

Only the Python standard library is needed. If openpyxl is installed an
Excel workbook of the results is written as well.

Usage:
    python water_hydraulics.py example_system.json [--out results]
"""
import argparse
import csv
import json
import math
import os
import sys

G = 9.81          # m/s2
RHO = 1000.0      # kg/m3
M_TO_BAR = RHO * G / 1e5

# material: Hazen-Williams C, absolute roughness k (mm), pressure wave speed a (m/s)
MATERIALS = {
    "DI":       (130, 0.10, 1100),
    "STEEL":    (120, 0.10, 1000),
    "GRP":      (140, 0.03, 500),
    "HDPE":     (140, 0.01, 300),
    "PVC":      (145, 0.0015, 400),
    "CONCRETE": (110, 0.50, 1100),
}

STANDARD_DN = [100, 150, 200, 250, 300, 350, 400, 450, 500, 600, 700, 800, 900,
               1000, 1100, 1200, 1400, 1500, 1600, 1800, 2000, 2200, 2400]
PN_CLASSES = [6, 10, 16, 25, 40, 50, 64, 100]
MOTOR_KW = [0.75, 1.1, 1.5, 2.2, 3, 4, 5.5, 7.5, 11, 15, 18.5, 22, 30, 37, 45, 55,
            75, 90, 110, 132, 160, 200, 250, 315, 355, 400, 450, 500, 560, 630,
            710, 800, 900, 1000, 1120, 1250, 1400, 1600, 1800, 2000, 2240, 2500,
            2800, 3150, 3550, 4000, 4500, 5000, 5600, 6300]

DEFAULTS = {
    "friction_method": "hazen-williams",   # or "darcy-weisbach"
    "water_temperature_c": 20.0,
    "minor_loss_factor": 0.10,         # minor losses as a fraction of friction (if no K given)
    "gravity_hours": 24.0,             # hours/day that gravity mains flow
    "economic_velocity": 1.2,          # m/s, target for "auto" sizing of pumped mains
    "max_velocity": 2.0,               # m/s, upper limit (also start point for gravity sizing)
    "min_velocity": 0.5,               # m/s, below this sediment may settle
    "min_pressure_head": 5.0,          # m, minimum pressure at junctions / delivery points
    "min_pipe_pressure": None,         # m, minimum pressure along mains (None = min_pressure_head)
    "pipe_cover": 0.0,                 # m, depth of pipe crown below ground in the profile
    "terminal_residual_head": 1.0,     # m, head kept above inlet level at tanks/sumps
    "station_losses": 3.0,             # m, pipework/valve losses inside a pump station
    "motor_margin": 1.15,              # motor rating >= shaft power x margin
    "surge_factor": 1.0,               # fraction of Joukowsky surge left after protection
    "pn_basis": "surge",               # "surge": PFA + PMA (=1.2 PFA) check incl. surge; "working": PFA only
    "npsh_margin": 1.0,                # m, NPSHa - NPSHr must exceed this
    "sump_retention_min": 10.0,        # min, sump volume >= pump flow x this
    "emergency_storage_hours": 0.0,    # h of outflow kept in tanks on top of balancing
    "electricity_tariff": 0.0,         # currency per kWh
    "currency": "USD",
    "standard_dn": STANDARD_DN,
    "pn_classes": PN_CLASSES,
}

PUMP_DEFAULTS = {"duty": 1, "standby": 1, "operating_hours": 22.0,
                 "efficiency": 0.75, "motor_efficiency": 0.94,
                 "suction_loss": 0.5, "npsh_required": None, "centerline_level": None}

NODE_TYPES = ("source", "intake", "sump", "wtp", "bps", "junction", "tank")


# ---------------------------------------------------------------- physics
def kinematic_viscosity(t_c):
    """Kinematic viscosity of water (m2/s) at t_c degC."""
    return 1.792e-6 / (1 + 0.0337 * t_c + 0.000221 * t_c ** 2)


def vapour_pressure_head(t_c):
    """Vapour pressure of water as metres of water (Buck equation)."""
    kpa = 0.61121 * math.exp((18.678 - t_c / 234.5) * (t_c / (257.14 + t_c)))
    return kpa * 1000 / (RHO * G)


def atmospheric_head(altitude_m):
    """Atmospheric pressure at altitude as metres of water."""
    kpa = 101.325 * (1 - 2.25577e-5 * altitude_m) ** 5.25588
    return kpa * 1000 / (RHO * G)


def friction_loss(q, d, length, method, c=130, k_mm=0.1, nu=1.004e-6):
    """Friction head loss (m) for flow q (m3/s) in a pipe of internal diameter d (m)."""
    if q <= 0 or length <= 0:
        return 0.0
    if method == "hazen-williams":
        return 10.67 * length * q ** 1.852 / (c ** 1.852 * d ** 4.8704)
    v = q / (math.pi * d * d / 4)
    re = v * d / nu
    if re < 2000:
        f = 64 / re
    else:  # Swamee-Jain explicit form of Colebrook-White
        f = 0.25 / math.log10(k_mm / 1000 / (3.7 * d) + 5.74 / re ** 0.9) ** 2
    return f * length / d * v * v / (2 * G)


def joukowsky(a, dv):
    """Surge head (m) for an instantaneous velocity change dv (m/s)."""
    return a * dv / G


def next_size(value, sizes):
    return next((s for s in sizes if s >= value - 1e-9), None)


def pn_segments(rows, col):
    """Group profile points into chainage ranges of equal PN; a section between two
    points takes the higher PN of its ends."""
    segs = []
    for r0, r1 in zip(rows, rows[1:]):
        a, b = r0[col], r1[col]
        pn = None if a is None or b is None else max(a, b)
        if segs and segs[-1][2] == pn:
            segs[-1][1] = r1[0]
        else:
            segs.append([r0[0], r1[0], pn])
    return segs


def max_pn(segs):
    return None if any(sg[2] is None for sg in segs) else max(sg[2] for sg in segs)


# ---------------------------------------------------------------- model
class Node:
    def __init__(self, d, settings):
        self.raw = d
        self.id = str(d["id"])
        self.type = d["type"].lower()
        if self.type not in NODE_TYPES:
            raise ValueError(f"node {self.id}: unknown type '{d['type']}' (use {', '.join(NODE_TYPES)})")
        self.name = d.get("name", self.id)
        self.min_level = d.get("min_level")          # lowest water level (source LWL, sump LWL, tank BWL)
        self.max_level = d.get("max_level")          # highest water level (TWL)
        self.elevation = d.get("elevation", self.min_level)   # ground level
        if self.elevation is None:
            raise ValueError(f"node {self.id}: give 'elevation' (ground level)")
        # level the incoming pipe must reach (inlet / weir level); default = TWL (or sump LWL at an intake)
        default_inlet = self.min_level if self.type == "intake" else (self.max_level or self.min_level)
        self.inlet_level = d.get("inlet_level", default_inlet)
        default_res = 0.0 if self.type == "intake" else settings["terminal_residual_head"]
        self.residual = d.get("residual_head", default_res)
        self.demand = float(d.get("demand_m3_day", 0.0))      # water drawn off at this node
        self.loss = float(d.get("process_loss", 0.0))         # e.g. 0.05 = 5 % WTP losses
        self.volume = d.get("volume_m3")                      # tank / sump / clear-water volume
        self.safe_yield = d.get("safe_yield_m3_day")
        p = d.get("pumps")
        self.pumps = {**PUMP_DEFAULTS, **p} if p else None
        if self.type != "junction" and self.min_level is None:
            raise ValueError(f"node {self.id}: free-surface node needs 'min_level'")
        # results
        self.inflow_daily = 0.0
        self.hgl = None
        self.excess_head = None
        self.station = None
        self.tank = None

    @property
    def is_free_surface(self):
        return self.type != "junction"


class Link:
    def __init__(self, d, settings):
        self.raw = d
        self.id = str(d["id"])
        self.frm = str(d["from"])
        self.to = str(d["to"])
        self.name = d.get("name", f"{self.frm} -> {self.to}")
        self.length = float(d["length"])
        self.auto = str(d.get("dn", "auto")).lower() == "auto"
        self.dn = None if self.auto else int(d["dn"])
        self.id_mm = d.get("id_mm")                   # internal diameter override
        self.material = d.get("material", "DI").upper()
        c, k, a = MATERIALS.get(self.material, MATERIALS["DI"])
        self.c = d.get("hw_c", c)
        self.k_mm = d.get("roughness_mm", k)
        self.wave_speed = d.get("wave_speed", a)
        self.minor_k = d.get("minor_k")              # sum of K values; else settings factor
        self.pn = d.get("pn")                        # installed / proposed PN (optional)
        self.profile = d.get("profile")              # [[chainage m, ground level m], ...]
        # results
        self.daily = self.hours = self.q = self.v = 0.0
        self.hf = self.hm = self.hloss = 0.0
        self.req_start = self.req_end = 0.0
        self.h_start = self.h_end = 0.0
        self.points = []
        self.gov_next = None

    @property
    def d(self):
        return (self.id_mm if self.id_mm else self.dn) / 1000.0

    @property
    def area(self):
        return math.pi * self.d ** 2 / 4


class System:
    def __init__(self, data):
        self.project = data.get("project", "Water supply system")
        self.settings = {**DEFAULTS, **data.get("settings", {})}
        self.settings["friction_method"] = self.settings["friction_method"].lower()
        self.nodes = {}
        for n in data["nodes"]:
            node = Node(n, self.settings)
            if node.id in self.nodes:
                raise ValueError(f"duplicate node id {node.id}")
            self.nodes[node.id] = node
        self.links = [Link(l, self.settings) for l in data["links"]]
        self.out = {nid: [] for nid in self.nodes}
        self.inc = {}
        for l in self.links:
            for end in (l.frm, l.to):
                if end not in self.nodes:
                    raise ValueError(f"link {l.id}: node '{end}' not defined")
            if l.to in self.inc:
                raise ValueError(f"node {l.to} is fed by two pipes ({self.inc[l.to].id}, {l.id}); "
                                 "only branched (tree) systems are supported")
            self.out[l.frm].append(l)
            self.inc[l.to] = l
        self.sources = [n for n in self.nodes.values() if n.type == "source"]
        if not self.sources:
            raise ValueError("no node of type 'source'")
        for n in self.nodes.values():
            if n.type == "source" and n.id in self.inc:
                raise ValueError(f"source {n.id} cannot have an incoming pipe")
            if n.type != "source" and n.id not in self.inc:
                raise ValueError(f"node {n.id} is not connected to any source")
            if n.type == "junction" and not self.out[n.id]:
                raise ValueError(f"junction {n.id} has no outgoing pipe")
            if n.pumps and not self.out[n.id]:
                raise ValueError(f"node {n.id} has pumps but no outgoing pipe")
        self.order = self._topological()
        self.warnings = []
        self.nu = kinematic_viscosity(self.settings["water_temperature_c"])

    def _topological(self):
        order, seen, stack = [], set(), [s.id for s in self.sources]
        while stack:
            nid = stack.pop()
            if nid in seen:
                raise ValueError(f"loop detected at node {nid}")
            seen.add(nid)
            order.append(nid)
            stack.extend(l.to for l in self.out[nid])
        missing = set(self.nodes) - seen
        if missing:
            raise ValueError(f"nodes not reachable from a source: {sorted(missing)}")
        return order

    def warn(self, obj, msg):
        self.warnings.append((obj, msg))

    # ------------------------------------------------------------ zones
    # A zone is the part of the system between free water surfaces: it starts at a
    # source / sump / tank (driven by pumps or by gravity) and runs through junctions
    # until the next free surface.
    def zone_starts(self):
        return [self.nodes[nid] for nid in self.order
                if self.nodes[nid].is_free_surface and self.out[nid]]

    def zone_links(self, start):
        links, stack = [], list(self.out[start.id])
        while stack:
            l = stack.pop()
            links.append(l)
            if self.nodes[l.to].type == "junction":
                stack.extend(self.out[l.to])
        return links

    def zone_ends(self, start):
        return [self.nodes[l.to] for l in self.zone_links(start)
                if self.nodes[l.to].is_free_surface]

    # ------------------------------------------------------------ solve
    def solve(self):
        self._flows()
        self._initial_diameters()
        for _ in range(500):
            self._headlosses()
            self._requirements()
            if not self._upsize_gravity():
                break
        self._forward()
        self._pipe_checks()
        self._stations()
        self._tanks()
        return self

    def _flows(self):
        def inflow(node):
            total = node.demand
            for l in self.out[node.id]:
                l.daily = inflow(self.nodes[l.to])
                total += l.daily
            if not 0 <= node.loss < 1:
                raise ValueError(f"node {node.id}: process_loss must be between 0 and 1")
            node.inflow_daily = total / (1 - node.loss)
            return node.inflow_daily

        for s in self.sources:
            inflow(s)
            if s.safe_yield is not None and s.inflow_daily > s.safe_yield:
                self.warn(s.id, f"abstraction {s.inflow_daily:,.0f} m3/day exceeds safe yield "
                                f"{s.safe_yield:,.0f} m3/day")

        for s in self.zone_starts():
            if s.pumps:
                hours = s.pumps["operating_hours"]
            else:
                # a gravity pipe that feeds an intake sump carries the intake pump rate
                pumped_ends = [e.pumps["operating_hours"] for e in self.zone_ends(s)
                               if e.type == "intake" and e.pumps]
                hours = min(pumped_ends) if pumped_ends else self.settings["gravity_hours"]
            for l in self.zone_links(s):
                l.hours = hours
                l.q = l.daily / (hours * 3600)

    def _initial_diameters(self):
        dns = self.settings["standard_dn"]
        for s in self.zone_starts():
            target = self.settings["economic_velocity"] if s.pumps else self.settings["max_velocity"]
            for l in self.zone_links(s):
                if l.auto:
                    d_min = math.sqrt(4 * l.q / (math.pi * target)) * 1000
                    l.dn = next_size(d_min, dns) or dns[-1]
                    l.id_mm = None

    def _headlosses(self):
        st = self.settings
        for l in self.links:
            l.v = l.q / l.area
            l.hf = friction_loss(l.q, l.d, l.length, st["friction_method"], l.c, l.k_mm, self.nu)
            if l.minor_k is not None:
                l.hm = l.minor_k * l.v ** 2 / (2 * G)
            else:
                l.hm = st["minor_loss_factor"] * l.hf
            l.hloss = l.hf + l.hm
            l.points = self._profile(l)

    def _profile(self, l):
        z0, z1 = self.nodes[l.frm].elevation, self.nodes[l.to].elevation
        pts = sorted((float(x), float(z)) for x, z in (l.profile or []) if 0 <= float(x) <= l.length)
        if not pts or pts[0][0] > 0:
            pts.insert(0, (0.0, z0))
        if pts[-1][0] < l.length:
            pts.append((l.length, z1))
        cover = self.settings["pipe_cover"]
        return [(x, z - cover) for x, z in pts]   # pipe crown level

    def _pressure_points(self, l):
        """Profile points where the minimum pressure applies: not at an end that
        sits on a free water surface (source, sump or tank connection)."""
        pts = l.points
        if self.nodes[l.frm].is_free_surface:
            pts = pts[1:]
        if self.nodes[l.to].is_free_surface:
            pts = pts[:-1]
        return pts

    def _requirements(self):
        """Backward pass: the HGL each pipe needs at its start so that every point
        downstream keeps the minimum pressure and every inlet is reached."""
        pmin = self.settings["min_pressure_head"]
        ppipe = self.pipe_pmin

        def req_end(node):
            if node.is_free_surface:
                return node.inlet_level + node.residual, None
            best, gov = node.elevation + pmin, None
            for c in self.out[node.id]:
                r = req_start(c)
                if r > best:
                    best, gov = r, c
            return best, gov

        def req_start(l):
            l.req_end, l.gov_next = req_end(self.nodes[l.to])
            r = l.req_end + l.hloss
            l.gov_point = None
            for x, z in self._pressure_points(l):
                # HGL at x = H_start - hloss * x / L must stay >= z + minimum pipe pressure
                rx = z + ppipe + l.hloss * x / l.length
                if rx > r + 1e-9:
                    r, l.gov_point = rx, (x, z)
            l.req_start = r
            return r

        for s in self.zone_starts():
            s.required_head = max(req_start(l) for l in self.out[s.id])
            if s.pumps:
                s.available_head = None
            else:
                s.available_head = s.min_level

    def governing(self, s):
        """Plain-language description of the point that sets the head needed at zone start s."""
        l = max(self.out[s.id], key=lambda x: x.req_start)
        while True:
            if l.gov_point is not None:
                x, z = l.gov_point
                return (f"pipe {l.id} at {x:,.0f} m (pipe crown {z:.1f} m + "
                        f"{self.pipe_pmin:.0f} m min. pressure)")
            end = self.nodes[l.to]
            if end.is_free_surface:
                return f"inlet of {end.id} (level {end.inlet_level:.1f} m + {end.residual:.1f} m)"
            if l.gov_next is None:
                return f"minimum pressure at junction {end.id}"
            l = l.gov_next

    def _critical_path(self, s):
        l = max(self.out[s.id], key=lambda x: x.req_start)
        path = [l]
        while l.gov_next is not None:
            l = l.gov_next
            path.append(l)
        return path

    def _upsize_gravity(self):
        dns = self.settings["standard_dn"]
        changed = False
        for s in self.zone_starts():
            if s.pumps or s.available_head >= s.required_head - 1e-6:
                continue
            cands = [l for l in self._critical_path(s)
                     if l.auto and next_size(l.dn + 1, dns) is not None]
            if cands:
                l = max(cands, key=lambda x: x.hloss / x.length)
                l.dn = next_size(l.dn + 1, dns)
                changed = True
        return changed

    def _forward(self):
        st = self.settings

        def propagate(l, h):
            l.h_start, l.h_end = h, h - l.hloss
            node = self.nodes[l.to]
            node.hgl = l.h_end
            if node.type == "junction":
                for c in self.out[node.id]:
                    propagate(c, l.h_end)
            else:
                node.excess_head = l.h_end - (node.inlet_level + node.residual)

        for s in self.zone_starts():
            if s.pumps:
                s.hgl = s.required_head                 # pumps deliver the critical requirement
                s.tdh = s.required_head - s.min_level + st["station_losses"]
                s.static_head = s.required_head - s.min_level
            else:
                s.hgl = s.available_head
                s.surplus = s.available_head - s.required_head
                if s.surplus < -1e-6:
                    self.warn(s.id, f"gravity zone from {s.name} is short of head by {-s.surplus:.1f} m "
                                    "- increase pipe sizes (or set dn 'auto') or add a booster pump")
            for l in self.out[s.id]:
                propagate(l, s.hgl)

    @property
    def pipe_pmin(self):
        v = self.settings["min_pipe_pressure"]
        return self.settings["min_pressure_head"] if v is None else v

    def _pipe_checks(self):
        st = self.settings
        pmin = self.pipe_pmin
        for s in self.zone_starts():
            if s.pumps:
                ends = self.zone_ends(s)
                static_level = max((e.max_level or e.inlet_level) for e in ends)
            else:
                static_level = s.max_level if s.max_level is not None else s.min_level
            for l in self.zone_links(s):
                surge = joukowsky(l.wave_speed, l.v) * st["surge_factor"]
                l.surge = surge
                rows = []
                for x, z in l.points:
                    hgl = l.h_start - l.hloss * x / l.length
                    p = hgl - z
                    working = max(p, static_level - z)
                    pn_w = next_size(working * M_TO_BAR, st["pn_classes"])
                    # EN 805: working pressure <= PFA (= PN) and working + surge <= PMA (= 1.2 PFA)
                    pn_s = next_size(max(working, (working + surge) / 1.2) * M_TO_BAR, st["pn_classes"])
                    rows.append((x, z, hgl, p, working, pn_w, pn_s))
                l.rows = rows
                checked = {x for x, _ in self._pressure_points(l)}
                l.p_min = min((r[3] for r in rows if r[0] in checked), default=None)
                l.p_max = max(r[3] for r in rows)
                l.p_work_max = max(r[4] for r in rows)
                l.transient_min = None if l.p_min is None else l.p_min - surge
                l.pn_segments_working = pn_segments(rows, 5)
                l.pn_segments_surge = pn_segments(rows, 6)
                l.pn_required_working = max_pn(l.pn_segments_working)
                l.pn_required_surge = max_pn(l.pn_segments_surge)
                if st["pn_basis"] == "working":
                    l.pn_segments, l.pn_required = l.pn_segments_working, l.pn_required_working
                else:
                    l.pn_segments, l.pn_required = l.pn_segments_surge, l.pn_required_surge

                tag = l.id
                if l.v > st["max_velocity"]:
                    self.warn(tag, f"velocity {l.v:.2f} m/s above {st['max_velocity']} m/s")
                if l.v < st["min_velocity"]:
                    self.warn(tag, f"velocity {l.v:.2f} m/s below {st['min_velocity']} m/s (sedimentation)")
                if l.p_min is None:
                    pass
                elif l.p_min < 0:
                    self.warn(tag, f"negative pressure {l.p_min:.1f} m along the pipe - HGL below ground")
                elif l.p_min < pmin - 1e-6:
                    self.warn(tag, f"minimum pressure {l.p_min:.1f} m below {pmin} m")
                if l.transient_min is not None and l.transient_min < -8 and l.length >= 200:
                    self.warn(tag, f"pump trip / valve closure can drop pressure to {l.transient_min:.0f} m "
                                   f"(surge {surge:.0f} m) - surge protection needed (surge vessel, air valves)")
                if l.pn_required is None:
                    self.warn(tag, "pressure exceeds the highest PN class - review surge protection / route")
                elif l.pn is not None and float(l.pn) < l.pn_required:
                    self.warn(tag, f"installed PN{l.pn} is below required PN{l.pn_required}")
                if l.auto and l.dn == st["standard_dn"][-1] and s.pumps is None and s.surplus < 0:
                    self.warn(tag, "reached the largest standard DN")

    def _stations(self):
        st = self.settings
        for n in self.nodes.values():
            if not n.pumps:
                continue
            p = n.pumps
            q_total = sum(l.q for l in self.out[n.id])
            daily = sum(l.daily for l in self.out[n.id])
            duty = int(p["duty"])
            q_pump = q_total / duty
            hyd_kw = RHO * G * q_pump * n.tdh / 1000
            shaft_kw = hyd_kw / p["efficiency"]
            motor_kw = next_size(shaft_kw * st["motor_margin"], MOTOR_KW)
            input_kw = shaft_kw / p["motor_efficiency"] * duty
            energy_day = input_kw * p["operating_hours"]
            r = {
                "q_total": q_total, "daily": daily, "q_pump": q_pump, "duty": duty,
                "standby": int(p["standby"]), "hours": p["operating_hours"],
                "suction_level": n.min_level, "delivery_hgl": n.hgl,
                "static_head": n.static_head, "tdh": n.tdh,
                "tdh_min": n.tdh - ((n.max_level or n.min_level) - n.min_level),
                "hyd_kw": hyd_kw, "shaft_kw": shaft_kw, "motor_kw": motor_kw,
                "input_kw": input_kw, "energy_day": energy_day,
                "kwh_per_m3": energy_day / daily if daily else 0.0,
                "cost_year": energy_day * 365 * st["electricity_tariff"],
                "npsha": None, "npshr": p["npsh_required"],
                "governed_by": self.governing(n),
            }
            if motor_kw is None:
                self.warn(n.id, f"shaft power {shaft_kw:,.0f} kW per pump is above the motor list - add duty pumps")
            if p["centerline_level"] is not None:
                r["npsha"] = (atmospheric_head(n.elevation) - vapour_pressure_head(st["water_temperature_c"])
                              + n.min_level - p["centerline_level"] - p["suction_loss"])
                if p["npsh_required"] is not None and r["npsha"] - p["npsh_required"] < st["npsh_margin"]:
                    self.warn(n.id, f"NPSH available {r['npsha']:.1f} m vs required {p['npsh_required']} m "
                                    f"- margin below {st['npsh_margin']} m (cavitation risk)")
            need = q_total * st["sump_retention_min"] * 60
            r["sump_needed"] = need
            r["sump_volume"] = n.volume
            if n.volume is not None and n.volume < need:
                self.warn(n.id, f"sump volume {n.volume:,.0f} m3 below {need:,.0f} m3 "
                                f"({st['sump_retention_min']:.0f} min at pump flow)")
            branch_heads = [l.req_start for l in self.out[n.id]]
            if max(branch_heads) - min(branch_heads) > 1:
                self.warn(n.id, "outgoing mains need different heads - flow control valves required on the "
                                "lower-head mains")
            n.station = r

    def _tanks(self):
        st = self.settings
        for n in self.nodes.values():
            if n.type != "tank":
                continue
            inc = self.inc[n.id]
            outflow = n.inflow_daily * (1 - n.loss)
            balancing = inc.daily * max(0.0, 1 - inc.hours / 24)
            emergency = outflow * st["emergency_storage_hours"] / 24
            needed = balancing + emergency
            r = {"inflow": inc.daily, "inflow_hours": inc.hours, "outflow": outflow,
                 "balancing": balancing, "emergency": emergency, "needed": needed,
                 "volume": n.volume,
                 "storage_hours": (n.volume / (outflow / 24)) if n.volume and outflow else None,
                 "arrival_hgl": n.hgl, "inlet": n.inlet_level, "excess": n.excess_head}
            if n.volume is not None and n.volume < needed:
                self.warn(n.id, f"volume {n.volume:,.0f} m3 below required {needed:,.0f} m3 "
                                "(balancing for pump-off hours + emergency)")
            n.tank = r
        for n in self.nodes.values():
            if n.is_free_surface and n.type != "source" and n.excess_head is not None and n.excess_head > 5:
                zone_start = self._zone_start_of(n)
                how = "flow control valve / PRV" if zone_start.pumps is None else "flow control valve"
                self.warn(n.id, f"arrives with {n.excess_head:.1f} m excess head - fit a {how} at the inlet"
                          + (" or reduce the pipe size" if zone_start.pumps is None else ""))

    def _zone_start_of(self, node):
        while True:
            node = self.nodes[self.inc[node.id].frm]
            if node.is_free_surface:
                return node

    # ------------------------------------------------------------ paths for the profile plot
    def paths(self):
        leaves = [n for n in self.nodes.values() if not self.out[n.id]]
        result = []
        for leaf in leaves:
            path, node = [], leaf
            while node.id in self.inc:
                l = self.inc[node.id]
                path.append(l)
                node = self.nodes[l.frm]
            result.append(list(reversed(path)))
        return result


# ---------------------------------------------------------------- reporting
def fmt(v, nd=1):
    if v is None:
        return "-"
    if isinstance(v, str):
        return v
    return f"{v:,.{nd}f}"


def table(headers, rows):
    cells = [[str(h) for h in headers]] + [[str(c) for c in r] for r in rows]
    w = [max(len(r[i]) for r in cells) for i in range(len(headers))]
    line = "  ".join("-" * x for x in w)
    out = ["  ".join(c.ljust(w[i]) if i == 0 else c.rjust(w[i]) for i, c in enumerate(cells[0])), line]
    for r in cells[1:]:
        out.append("  ".join(c.ljust(w[i]) if i == 0 else c.rjust(w[i]) for i, c in enumerate(r)))
    return "\n".join(out)


def pn_text(l, segs=None):
    parts = []
    for x0, x1, pn in (l.pn_segments if segs is None else segs):
        parts.append(f"{'PN' + str(pn) if pn else '>PN max'} ({x0:,.0f}-{x1:,.0f} m)")
    return ", ".join(parts)


def pipe_rows(sys_):
    rows = []
    for l in sys_.links:
        rows.append({
            "Pipe": l.id, "Name": l.name, "From": l.frm, "To": l.to,
            "Material": l.material, "DN": l.dn, "ID (mm)": round(l.d * 1000, 1),
            "Sized": "auto" if l.auto else "given", "Length (m)": l.length,
            "Hours/day": l.hours, "Q (m3/day)": l.daily, "Q (m3/h)": l.q * 3600,
            "Q (l/s)": l.q * 1000, "V (m/s)": l.v, "Friction loss (m)": l.hf,
            "Minor loss (m)": l.hm, "Total loss (m)": l.hloss,
            "Gradient (m/km)": l.hloss / l.length * 1000,
            "HGL start (m)": l.h_start, "HGL end (m)": l.h_end,
            "Min pressure (m)": l.p_min, "Max pressure (m)": l.p_max,
            "Surge dH (m)": l.surge, "Transient min (m)": l.transient_min,
            "PN required": l.pn_required, "PN installed": l.pn,
            "PN working only": l.pn_required_working, "PN incl. surge": l.pn_required_surge,
            "PN by chainage (working)": pn_text(l, l.pn_segments_working),
            "PN by chainage (incl. surge)": pn_text(l, l.pn_segments_surge),
        })
    return rows


def profile_rows(sys_):
    rows = []
    for l in sys_.links:
        for x, z, hgl, p, working, pn_w, pn_s in l.rows:
            rows.append({"Pipe": l.id, "From": l.frm, "To": l.to, "Chainage in pipe (m)": x,
                         "Ground (m)": z, "HGL (m)": hgl, "Pressure (m)": p,
                         "Working pressure incl. static (m)": working,
                         "PN working": pn_w, "PN incl. surge": pn_s})
    return rows


def station_rows(sys_):
    rows = []
    for n in sys_.nodes.values():
        r = n.station
        if not r:
            continue
        rows.append({
            "Station": n.id, "Name": n.name, "Type": n.type.upper(),
            "Pumps (duty+standby)": f"{r['duty']}w+{r['standby']}s", "Hours/day": r["hours"],
            "Q total (m3/day)": r["daily"], "Q total (m3/h)": r["q_total"] * 3600,
            "Q per pump (m3/h)": r["q_pump"] * 3600, "Q per pump (l/s)": r["q_pump"] * 1000,
            "Suction LWL (m)": r["suction_level"], "Delivery HGL (m)": r["delivery_hgl"],
            "Static+friction (m)": r["static_head"], "TDH design (m)": r["tdh"],
            "TDH at sump TWL (m)": r["tdh_min"], "Hydraulic kW/pump": r["hyd_kw"],
            "Shaft kW/pump": r["shaft_kw"], "Motor kW/pump": r["motor_kw"],
            "Input kW (all duty)": r["input_kw"], "Energy (kWh/day)": r["energy_day"],
            "kWh/m3": r["kwh_per_m3"], "Energy cost/yr": r["cost_year"],
            "NPSHa (m)": r["npsha"], "NPSHr (m)": r["npshr"],
            "Sump volume (m3)": r["sump_volume"], "Sump needed (m3)": r["sump_needed"],
            "Head governed by": r["governed_by"],
        })
    return rows


def tank_rows(sys_):
    rows = []
    for n in sys_.nodes.values():
        r = n.tank
        if not r:
            continue
        rows.append({
            "Tank": n.id, "Name": n.name, "BWL (m)": n.min_level, "TWL (m)": n.max_level,
            "Inlet level (m)": r["inlet"], "Arrival HGL (m)": r["arrival_hgl"],
            "Excess head (m)": r["excess"], "Inflow (m3/day)": r["inflow"],
            "Inflow hours": r["inflow_hours"], "Outflow (m3/day)": r["outflow"],
            "Local demand (m3/day)": n.demand, "Balancing (m3)": r["balancing"],
            "Emergency (m3)": r["emergency"], "Required (m3)": r["needed"],
            "Volume (m3)": r["volume"], "Storage (h of outflow)": r["storage_hours"],
        })
    return rows


def gravity_rows(sys_):
    rows = []
    for s in sys_.zone_starts():
        if s.pumps:
            continue
        rows.append({"Zone start": s.id, "Name": s.name, "Available head (m)": s.available_head,
                     "Required head (m)": s.required_head, "Surplus (m)": s.surplus,
                     "Status": "OK" if s.surplus >= -1e-6 else "SHORT OF HEAD",
                     "Governed by": sys_.governing(s)})
    return rows


def text_report(sys_):
    st = sys_.settings
    L = []
    L.append("=" * 100)
    L.append(f"HYDRAULIC ANALYSIS - {sys_.project}")
    L.append("=" * 100)
    L.append(f"Friction: {st['friction_method']}, minor losses "
             f"{'K values' if any(l.minor_k is not None for l in sys_.links) else ''}"
             f"{st['minor_loss_factor']:.0%} of friction where no K given, water {st['water_temperature_c']} degC")
    L.append(f"Min pressure {st['min_pressure_head']} m at junctions, {sys_.pipe_pmin} m along mains "
             f"(pipe crown {st['pipe_cover']} m below ground), station losses {st['station_losses']} m, "
             f"surge factor {st['surge_factor']} x Joukowsky, PN check EN 805 (PMA = 1.2 x PFA)")
    L.append("")
    total = sum(s.inflow_daily for s in sys_.sources)
    delivered = sum(n.demand for n in sys_.nodes.values())
    L.append(f"Raw water abstraction: {total:,.0f} m3/day    Water delivered to demand nodes: {delivered:,.0f} m3/day")
    L.append("")

    L.append("PIPES")
    L.append(table(["Pipe", "From -> To", "DN", "L (m)", "h/d", "Q (m3/d)", "Q (l/s)", "V (m/s)",
                    "Loss (m)", "m/km", "HGL in", "HGL out", "Pmin", "Pmax", "Surge", "PN"],
                   [[l.id, f"{l.frm} -> {l.to}", f"{l.dn}{'*' if l.auto else ''}", fmt(l.length, 0),
                     fmt(l.hours, 0), fmt(l.daily, 0), fmt(l.q * 1000, 0), fmt(l.v, 2), fmt(l.hloss, 1),
                     fmt(l.hloss / l.length * 1000, 2), fmt(l.h_start), fmt(l.h_end), fmt(l.p_min),
                     fmt(l.p_max), fmt(l.surge, 0), f"PN{l.pn_required}" if l.pn_required else ">max"]
                    for l in sys_.links]))
    L.append(f"  * = diameter chosen by the program; pressures in m of water; PN column basis: {st['pn_basis']}")
    L.append("  PN by chainage along each pipe - working pressure only / including surge:")
    for l in sys_.links:
        L.append(f"    {l.id}: {pn_text(l, l.pn_segments_working)}")
        L.append(f"    {' ' * len(l.id)}  {pn_text(l, l.pn_segments_surge)}  (incl. surge {l.surge:.0f} m)")
    L.append("")

    L.append("PUMP STATIONS")
    rows = []
    for n in sys_.nodes.values():
        r = n.station
        if r:
            rows.append([n.id, n.type.upper(), f"{r['duty']}w+{r['standby']}s", fmt(r["hours"], 0),
                         fmt(r["q_pump"] * 3600, 0), fmt(r["q_pump"] * 1000, 0), fmt(r["suction_level"]),
                         fmt(r["delivery_hgl"]), fmt(r["tdh"]), fmt(r["shaft_kw"], 0), fmt(r["motor_kw"], 0),
                         fmt(r["energy_day"], 0), fmt(r["kwh_per_m3"], 2), fmt(r["npsha"])])
    L.append(table(["Station", "Type", "Pumps", "h/d", "Q/pump m3/h", "l/s", "Suction LWL",
                    "Delivery HGL", "TDH (m)", "Shaft kW", "Motor kW", "kWh/day", "kWh/m3", "NPSHa"], rows))
    for n in sys_.nodes.values():
        if n.station:
            L.append(f"  {n.id}: head set by {n.station['governed_by']}")
    L.append("")

    g = gravity_rows(sys_)
    if g:
        L.append("GRAVITY ZONES")
        L.append(table(["Start", "Available (m)", "Required (m)", "Surplus (m)", "Status", "Governed by"],
                       [[r["Zone start"], fmt(r["Available head (m)"]), fmt(r["Required head (m)"]),
                         fmt(r["Surplus (m)"]), r["Status"], r["Governed by"]] for r in g]))
        L.append("")

    L.append("TANKS")
    rows = []
    for n in sys_.nodes.values():
        r = n.tank
        if r:
            rows.append([n.id, fmt(n.min_level), fmt(n.max_level), fmt(r["arrival_hgl"]), fmt(r["excess"]),
                         fmt(r["inflow"], 0), fmt(r["inflow_hours"], 0), fmt(r["needed"], 0),
                         fmt(r["volume"], 0), fmt(r["storage_hours"])])
    L.append(table(["Tank", "BWL", "TWL", "Arrival HGL", "Excess head", "Inflow m3/d", "h/d",
                    "Required m3", "Volume m3", "Storage h"], rows))
    L.append("")

    L.append("WARNINGS / DESIGN NOTES")
    if sys_.warnings:
        for obj, msg in sys_.warnings:
            L.append(f"  [{obj}] {msg}")
    else:
        L.append("  none")
    return "\n".join(L)


def write_csv(path, rows):
    if not rows:
        return
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        for r in rows:
            w.writerow({k: (round(v, 3) if isinstance(v, float) else v) for k, v in r.items()})


def write_excel(path, sys_, report, extra_sheets=()):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E78")
    sheets = [("Pipes", pipe_rows(sys_)), ("Pump stations", station_rows(sys_)),
              ("Tanks", tank_rows(sys_)), ("Gravity zones", gravity_rows(sys_)),
              ("Profile points", profile_rows(sys_))] + [(t[:31], r) for t, r in extra_sheets]
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Hydraulic analysis - {sys_.project}"
    ws["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(report.splitlines(), start=3):
        ws.cell(row=i, column=1, value=line).font = Font(name="Consolas", size=9)
    ws.column_dimensions["A"].width = 160
    for title, rows in sheets:
        if not rows:
            continue
        ws = wb.create_sheet(title)
        headers = list(rows[0])
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r, row in enumerate(rows, 2):
            for c, h in enumerate(headers, 1):
                v = row[h]
                cell = ws.cell(row=r, column=c, value=v)
                if isinstance(v, float):
                    cell.number_format = "#,##0.00" if abs(v) < 100 else "#,##0"
        for c, h in enumerate(headers, 1):
            longest = max([len(str(h)) // 2] + [len(str(row[h])) for row in rows])
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, longest + 2), 60)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"
    ws = wb.create_sheet("Warnings")
    ws.append(["Item", "Warning / design note"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for obj, msg in sys_.warnings:
        ws.append([obj, msg])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 120
    wb.save(path)
    return True


def write_profile_svg(path, sys_, links):
    """Longitudinal section: ground line and HGL from the source to one end tank."""
    ground, hgl, marks = [], [], []
    x0 = 0.0
    for i, l in enumerate(links):
        frm = sys_.nodes[l.frm]
        if frm.is_free_surface:
            level = frm.min_level
            if hgl:
                hgl.append((x0, level))          # water drops into the sump/tank
            if frm.pumps:
                hgl.append((x0, level))
                marks.append((x0, frm, f"{frm.id} (TDH {frm.tdh:.0f} m)"))
            else:
                marks.append((x0, frm, frm.id))
        elif i:
            marks.append((x0, frm, frm.id))
        for x, z in l.points:
            ground.append((x0 + x, z + sys_.settings["pipe_cover"]))
        for x, z, h, *_ in l.rows:
            hgl.append((x0 + x, h))
        x0 += l.length
    end = sys_.nodes[links[-1].to]
    marks.append((x0, end, end.id))
    if end.is_free_surface:
        hgl.append((x0, end.max_level or end.min_level))

    W, H, ml, mr, mt, mb = 1200, 560, 80, 30, 50, 60
    xs = [p[0] for p in ground]
    ys = [p[1] for p in ground] + [p[1] for p in hgl]
    xmax = max(xs) or 1
    ymin, ymax = min(ys), max(ys)
    pad = (ymax - ymin) * 0.08 or 5
    ymin, ymax = ymin - pad, ymax + pad

    def X(x):
        return ml + x / xmax * (W - ml - mr)

    def Y(y):
        return mt + (ymax - y) / (ymax - ymin) * (H - mt - mb)

    def poly(pts):
        return " ".join(f"{X(x):.1f},{Y(y):.1f}" for x, y in pts)

    o = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}" '
         'font-family="Arial, sans-serif" font-size="12">',
         f'<rect width="{W}" height="{H}" fill="#ffffff"/>',
         f'<text x="{ml}" y="24" font-size="16" font-weight="bold" fill="#1a1a1a">'
         f'Hydraulic profile: {links[0].frm} to {end.id}</text>']
    # grid
    step_y = nice_step((ymax - ymin) / 8)
    y = math.ceil(ymin / step_y) * step_y
    while y <= ymax:
        o.append(f'<line x1="{ml}" x2="{W - mr}" y1="{Y(y):.1f}" y2="{Y(y):.1f}" stroke="#e5e5e5"/>')
        o.append(f'<text x="{ml - 8}" y="{Y(y) + 4:.1f}" text-anchor="end" fill="#555">{y:,.0f}</text>')
        y += step_y
    step_x = nice_step(xmax / 10)
    x = 0.0
    while x <= xmax + 1e-6:
        o.append(f'<text x="{X(x):.1f}" y="{H - mb + 18}" text-anchor="middle" fill="#555">{x / 1000:,.1f}</text>')
        x += step_x
    o.append(f'<text x="{(W + ml) / 2}" y="{H - 18}" text-anchor="middle" fill="#333">Chainage (km)</text>')
    o.append(f'<text transform="translate(20,{(H + mt - mb) / 2}) rotate(-90)" text-anchor="middle" '
             'fill="#333">Level (m)</text>')
    # ground (filled) and HGL
    fill = [(ground[0][0], ymin)] + ground + [(ground[-1][0], ymin)]
    o.append(f'<polygon points="{poly(fill)}" fill="#d9c9a8" fill-opacity="0.55" stroke="none"/>')
    o.append(f'<polyline points="{poly(ground)}" fill="none" stroke="#8a6d3b" stroke-width="1.5"/>')
    o.append(f'<polyline points="{poly(hgl)}" fill="none" stroke="#1f6fd1" stroke-width="2.2"/>')
    # stations - labels stepped down when stations are close together
    last_x, level = -1e9, 0
    for x, node, label in marks:
        o.append(f'<line x1="{X(x):.1f}" x2="{X(x):.1f}" y1="{mt}" y2="{H - mb}" stroke="#999" '
                 'stroke-dasharray="4 4"/>')
        level = level + 1 if X(x) - last_x < 170 else 0
        last_x = X(x)
        anchor = "start" if x < xmax * 0.85 else "end"
        dx = 4 if anchor == "start" else -4
        o.append(f'<text x="{X(x) + dx:.1f}" y="{mt + 14 + 15 * (level % 6)}" text-anchor="{anchor}" '
                 f'fill="#1a1a1a">{escape(label)}</text>')
        if node.is_free_surface:
            o.append(f'<circle cx="{X(x):.1f}" cy="{Y(node.min_level):.1f}" r="3.5" fill="#1a1a1a"/>')
    # legend
    lx = W - mr - 220
    o.append(f'<line x1="{lx}" x2="{lx + 30}" y1="{mt - 16}" y2="{mt - 16}" stroke="#1f6fd1" stroke-width="2.2"/>')
    o.append(f'<text x="{lx + 36}" y="{mt - 12}" fill="#1a1a1a">HGL</text>')
    o.append(f'<rect x="{lx + 90}" y="{mt - 22}" width="30" height="12" fill="#d9c9a8" stroke="#8a6d3b"/>')
    o.append(f'<text x="{lx + 126}" y="{mt - 12}" fill="#1a1a1a">Ground</text>')
    o.append("</svg>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(o))


def nice_step(raw):
    if raw <= 0:
        return 1
    mag = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if raw <= m * mag:
            return m * mag
    return 10 * mag


def escape(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------------------------------------------------------------- CLI
def run(input_path, out_dir, svg=True, excel=True, quiet=False):
    with open(input_path, encoding="utf-8") as f:
        data = json.load(f)
    s = System(data).solve()
    report = text_report(s)
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "report.txt"), "w", encoding="utf-8") as f:
        f.write(report + "\n")
    write_csv(os.path.join(out_dir, "pipes.csv"), pipe_rows(s))
    write_csv(os.path.join(out_dir, "pump_stations.csv"), station_rows(s))
    write_csv(os.path.join(out_dir, "tanks.csv"), tank_rows(s))
    write_csv(os.path.join(out_dir, "profile_points.csv"), profile_rows(s))
    written = ["report.txt", "pipes.csv", "pump_stations.csv", "tanks.csv", "profile_points.csv"]
    if svg:
        for path in s.paths():
            name = f"profile_{path[0].frm}_to_{path[-1].to}.svg"
            write_profile_svg(os.path.join(out_dir, name), s, path)
            written.append(name)
    if excel:
        if write_excel(os.path.join(out_dir, "hydraulic_results.xlsx"), s, report):
            written.append("hydraulic_results.xlsx")
        elif not quiet:
            print("(openpyxl not installed - Excel output skipped: pip install openpyxl)")
    if not quiet:
        print(report)
        print(f"\nResults written to {out_dir}: {', '.join(written)}")
    return s


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="system description (JSON)")
    ap.add_argument("--out", default="results", help="output folder (default: results)")
    ap.add_argument("--no-svg", action="store_true", help="skip hydraulic profile drawings")
    ap.add_argument("--no-excel", action="store_true", help="skip the Excel workbook")
    a = ap.parse_args(argv)
    try:
        run(a.input, a.out, svg=not a.no_svg, excel=not a.no_excel)
    except (ValueError, KeyError) as e:
        sys.exit(f"Input error: {e}")


if __name__ == "__main__":
    main()
