#!/usr/bin/env python3
"""Copy the levels (and demands) typed into the Excel zone calculator back into tunduma_config.json.

    python levels_from_excel.py results/Tunduma_Zone_Calculator.xlsx [--config tunduma_config.json]

Then run:  python run_analysis.py ; python make_report_pdf.py ; python make_explainer_pdf.py ;
           python make_zone_workbook.py   (all reports use the same levels)

Sheet 'Levels' columns: ID | facility | role | ground | floor | MIN level | MAX level | inlet | capacity
Sheet 'Demands': the six tank demands and the WTP losses.
"""
import argparse
import json
import shutil

import openpyxl

LEVEL_FIELDS = (("elevation", 4), ("floor", 5), ("min_level", 6), ("max_level", 7), ("inlet_level", 8),
                ("volume_m3", 9))
DEMAND_LABELS = {"Uhuru Park GBR – town supply": "UHURU_PARK_GBR", "Existing Uhuru GSR 01": "UHURU_GSR01",
                 "Existing Tunduma ESR": "TUNDUMA_ESR", "Existing Ipito GSR 02": "IPITO_GSR02",
                 "Existing Makambini GSR 03": "MAKAMBINI_GSR03", "Existing Chapwa GSR 04": "CHAPWA_GSR04"}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--config", default="tunduma_config.json")
    a = ap.parse_args(argv)

    wb = openpyxl.load_workbook(a.workbook, data_only=True)
    with open(a.config, encoding="utf-8") as f:
        cfg = json.load(f)
    fac = {f["id"]: f for f in cfg["facilities"]}
    changes = []

    def setv(f, key, value):
        old = f.get(key)
        if value is None:
            if key in ("inlet_level", "floor") and key in f:
                del f[key]
                changes.append(f"{f['id']}.{key}: {old} -> (blank)")
            return
        value = round(float(value), 3)
        if old is None or abs(float(old) - value) > 1e-6:
            f[key] = value
            changes.append(f"{f['id']}.{key}: {old} -> {value}")

    ws = wb["Levels"]
    for row in ws.iter_rows(min_row=5, values_only=True):
        fid = row[0]
        if fid not in fac:
            continue
        lo, hi = row[5], row[6]
        if lo is None or hi is None or float(hi) <= float(lo):
            raise SystemExit(f"Levels sheet: {fid} needs MIN and MAX levels with MAX above MIN")
        for key, col in LEVEL_FIELDS:
            setv(fac[fid], key, row[col - 1])

    ws = wb["Demands"]
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[0] in DEMAND_LABELS:
            setv(fac[DEMAND_LABELS[row[0]]], "demand_m3_day", row[1])
        elif row[0] == "WTP + intake-to-WTP losses":
            setv(fac["MOMBA_WTP"], "process_loss", row[1])

    if not changes:
        print("No changes – the config already matches the workbook.")
        return
    shutil.copy(a.config, a.config + ".bak")
    with open(a.config, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"Updated {a.config} (backup {a.config}.bak):")
    for c in changes:
        print("  " + c)


if __name__ == "__main__":
    main()
